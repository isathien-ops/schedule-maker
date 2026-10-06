"""Выгрузка расписания в Excel.

Каждая функция ``…Workbook`` возвращает книгу openpyxl (Workbook); вызывающий код сам сохраняет
её куда нужно. Дни и время уроков берутся из сетки на вкладке «Настройки», выходные тоже.

Виды выгрузки:
* ``coursesWorkbook`` — неделя, столбец на каждый курс (предмет и преподаватель);
* ``teachersWorkbook`` — неделя, столбец на каждого преподавателя (какой курс он ведёт);
* ``calendarWorkbook`` — все уроки по датам учебного календаря, по строке на урок
  (уроки по датам разворачивает ``generateCalendarEvents``);
* ``variantsWorkbook`` — варианты одного этапа с «Предпросмотра»: сравнение и неделя каждого;
* ``teacherListWorkbook`` — преподаватели: их курсы («ведёт», «может вести», «запрещено») и отметки времени.

``answer`` везде — принятое расписание (answer.json): курс -> неделя -> день -> ячейка.
Неверные даты курсов (их нельзя разобрать) дают ``CourseDatesError`` — сервер показывает её
человеку понятным текстом; другие ошибки под неё не маскируются.

Общие уроки («линейка идёт вместе с Потоком N», модуль joint): у копии в answer те же уроки, что
у источника. В книгах курсов, календаря и вариантов они подписаны (``jointNotes``: у копии «вместе
с Потоком 1», у источника «вместе с Потоком 2»), а преподавателю общий урок — одна строка
(«Потоки 1 и 2») и один урок в неделю, копия — не отдельный курс в его списке.

Зависимости: model, grid, courses (подпись курса, отношение преподавателя к курсу), joint (курсы-копии
и их источники), penalties (ключи строк своих правил), variants (строки «Кто ведёт» листа сравнения),
school_defaults (даты календаря по умолчанию).
"""

import datetime
import hashlib
import re

import natsort
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as letter

from src.modules.functions.courses import copyFollowers, copyRoot, courseLabel, lineOf, teacherCourseState
from src.modules.functions.grid import DAYS, dayGrid
from src.modules.functions.joint import jointCopies
from src.modules.functions.model import (
    allLessons, cannotSlots, courseGroups, courseNames, courseSubject, groupsByName, isLesson, lessonEntries, possibleSlots
)
from src.modules.functions.penalties import customId
from src.modules.functions.school_defaults import CALENDAR_END, CALENDAR_START
from src.modules.functions.variants import teacherCourses, unacceptable
from src.modules.translate import fill

# Excel ограничивает имя листа 31 символом
SHEET_NAME_LIMIT = 31

# Отметки времени преподавателя на листах этапов: вид -> (ключ текста в ru.hjson, цвет как на
# странице преподавателя). Виды названы по смыслу: "ok" — «удобно», "may" — «может»,
# "cannot" — «не может» (как они хранятся в данных проекта — см. model.teacherAvailability)
MARKS = {
    "ok": ("menu.main.tab.teachers.free", "D1FADF"),
    "may": ("menu.main.tab.teachers.possible", "FEF0C7"),
    "cannot": ("menu.main.tab.teachers.busy", "FEE4E2"),
}

# Заливка ячейки «Кто ведёт» на листе сравнения, если вариант сменит преподавателя принятого курса
# (как выделенная ячейка таблицы на «Предпросмотре»)
CHANGED_FILL = "FEF0C7"

# Ширина столбцов вариантов на листе сравнения: под числа правил хватает 16, а с блоком «Кто ведёт»
# в ячейке полные имена («Калуцкова Анжелика» и «Сейчас ведёт: …») — тогда шире, чтобы имя не рвалось
VARIANT_WIDTH, TEACHER_WIDTH = 16, 26

# Цвет текста ячейки урока на листе варианта (variantSheet): красный — у урока есть проблема,
# тёмно-жёлтый — только некритичные подписи (те же цвета, что --bad и --warn в app.css)
PROBLEM_COLOR = "B42318"
WARNING_COLOR = "B54708"


class CourseDatesError(ValueError):
    """Дату курса (или календаря) нельзя разобрать: книга по датам не строится."""


def color(text):
    """Цвет заливки "RRGGBB" для текста (обычно названия предмета)."""
    # Стабильный пастельный цвет по тексту, чтобы у одного предмета везде был один цвет:
    # хэш даёт случайный цвет, а усреднение каждого канала с 255 делает его светлым
    number = int.from_bytes(hashlib.sha256(text.encode()).digest(), byteorder="big") % 0x1000000
    hexnum = f"{number:06x}"

    return "".join(f"{int((int(hexnum[i:i + 2], 16) + 255) / 2 + 0.5):02x}" for i in range(0, 6, 2))


def solidFill(rgb):
    """Сплошная заливка ячейки цветом "RRGGBB"."""
    return PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")


def sheetName(text):
    """Имя листа, обрезанное до допустимых в Excel 31 символа."""
    return text[:SHEET_NAME_LIMIT]


def gridRows(settings):
    """(сетка, дни, в которых есть уроки)."""
    grid = dayGrid(settings)
    days = [day for day in range(DAYS) if grid[day]]

    return grid, days


def lessonText(cell, translate, prefix=""):
    """Подпись урока: «Химия — Иванова» (без преподавателя — «Химия — без преподавателя», текст
    web.classes.no_teacher), с необязательной приставкой."""
    return f"{prefix}{cell['subject']} — {', '.join(cell.get('teachers', [])) or translate('web.classes.no_teacher')}"


def withNote(text, note):
    """Текст с пометкой в скобках: «Химия — Иванова (вместе с Потоком 1)»; без пометки — как есть."""
    return f"{text} ({note})" if note else text


# ---------------------------------------------------------------- общие уроки потоков

def streamsText(numbers):
    """Номера потоков для подписи: «2», «2 и 3», «1, 2 и 3» (как andList в подписях общих уроков на странице)."""
    texts = [str(number) for number in numbers]

    return f"{', '.join(texts[:-1])} и {texts[-1]}" if len(texts) > 1 else texts[0]


def jointStreams(settings, copies):
    """Потоки общих уроков: {курс-источник: [поток источника, потоки его копий…]} по возрастанию.

    Только источники, с которыми идёт хотя бы одна копия; ``copies`` — joint.jointCopies.
    Источник всегда в более раннем потоке, поэтому его номер первый.
    """
    groups = groupsByName(settings)

    return {source: sorted(groups[name]["stream_id"] for name in [source] + followers) for source, followers in copyFollowers(copies).items()}


def jointText(translate, numbers):
    """«вместе с Потоком 1» или «вместе с Потоками 2 и 3» (web.lesson.joint / joint_many)."""
    key = "web.lesson.joint_many" if len(numbers) > 1 else "web.lesson.joint"

    return fill(translate(key), number=streamsText(numbers), numbers=streamsText(numbers))


def jointNotes(settings, translate):
    """Пометки общих уроков: {курс: текст}, как подписи карточек на странице.

    У копии — её поток-источник («вместе с Потоком 1»), у источника — потоки его копий («вместе
    с Потоком 2»). У остальных курсов пометки нет: уроки копии в answer настоящие, и без пометки
    в книге не видно, что это те же уроки, а не вторые такие же.
    """
    copies = jointCopies(settings)
    streams = jointStreams(settings, copies)
    notes = {name: jointText(translate, streams[source][:1]) for name, source in copies.items()}
    notes.update({source: jointText(translate, numbers[1:]) for source, numbers in streams.items()})

    return notes


def weekSheet(page, settings, translate):
    """Рисует на листе столбцы «День / # / Время»; возвращает {(день, урок): строка} и строки-разделители.

    Под каждым днём оставляется пустая строка-разделитель (её потом красит ``finish``),
    ячейка дня объединяется на все его уроки.
    """
    grid, days = gridRows(settings)

    page["A1"] = translate("menu.main.tab.export.day")
    page["B1"] = "#"
    page["C1"] = translate("menu.main.tab.export.time")

    rows = {}
    separators = []
    row = 2

    for day in days:
        first = row

        for lesson, time in enumerate(grid[day]):
            page[f"B{row}"] = lesson + 1
            page[f"C{row}"] = time
            rows[(day, lesson)] = row
            row += 1

        page.merge_cells(f"A{first}:A{row - 1}")
        page[f"A{first}"] = translate(f"abbreviate.day.{day}")

        separators.append(row)
        row += 1

    return rows, separators


def finish(page, separators):
    """Оформление листа: закрепление шапки, выравнивание, ширина столбцов, серые разделители дней."""
    # Строка заголовков и столбцы день / # / время остаются на виду при прокрутке широкого листа
    page.freeze_panes = "D2"
    page.row_dimensions[1].height = 20

    # Ширина столбца — по самой длинной строке текста в нём (не больше 60 символов)
    for col in page.columns:
        width = 0

        for cell in col:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            if cell.value is not None:
                width = max(width, max(len(part) for part in str(cell.value).split("\n")))

        page.column_dimensions[letter(col[0].column)].width = min(width + 2, 60)

    # Разделители между днями: тонкие серые полосы
    for line in separators:
        page.row_dimensions[line].height = 3

        for col in range(1, page.max_column + 1):
            page.cell(row=line, column=col).fill = solidFill("D3D3D3")


# ---------------------------------------------------------------- книги по принятому расписанию

def coursesWorkbook(settings, answer, translate):
    """Столбец на каждый курс: предмет и преподаватель на каждом уроке.

    Курсы расписания, которых больше нет в настройках, не выгружаются; уроки вне текущей
    сетки (например, после её правки) пропускаются. Уроки общих курсов подписаны (``jointNotes``):
    «Математика — Иванова (вместе с Потоком 1)».
    """
    book = Workbook()
    page = book.active
    page.title = translate("menu.main.tab.export.schedule")

    rows, separators = weekSheet(page, settings, translate)
    courses = courseNames(settings)
    notes = jointNotes(settings, translate)

    # Курсы начинаются со столбца D (после «День / # / Время»)
    for col, course in enumerate(courses, start=4):
        page[f"{letter(col)}1"] = course

    for course, day, lesson, cell in allLessons(answer):
        if course in courses and (day, lesson) in rows:
            text = withNote(lessonText(cell, translate), notes.get(course))
            target = page.cell(row=rows[(day, lesson)], column=courses.index(course) + 4, value=text)
            target.fill = solidFill(color(cell["subject"]))

    finish(page, separators)

    return book


def teacherColumnLabel(group, course):
    """Подпись урока в столбце преподавателя: «Поток 2 (с 23.11), 10 класс, Химия» — курс и дата его старта.

    Подпись собирается из полей курса (courses.courseLabel), а не разбором его названия. Курса
    ``course`` уже нет в настройках (``group`` пустой) — пишется просто его название.
    """
    start = group.get("start_date") or ""
    since = f" (с {start[8:10]}.{start[5:7]})" if len(start) == 10 else ""

    return courseLabel(group, since) if group else course


def teachersWorkbook(settings, answer, translate):
    """Столбец на каждого преподавателя: курс на каждом его уроке.

    Если в одном слоте у преподавателя несколько курсов (разные даты), они пишутся
    через перенос строки. Общий урок источника и его копий — это один урок: одна строка с подписью
    источника и пометкой «Потоки 1 и 2», а не несколько строк, как при накладке. Уроки, которые
    никто не ведёт, получают собственный столбец, чтобы не пропасть молча.
    """
    lessons = list(allLessons(answer))
    copies = jointCopies(settings)
    # Пометка общего урока у источника: «Потоки 1 и 2» (web.lesson.joint_streams, как на странице)
    shared = {source: fill(translate("web.lesson.joint_streams"), numbers=streamsText(numbers))
              for source, numbers in jointStreams(settings, copies).items()}
    # Все преподаватели из расписания, в «естественном» порядке сортировки
    teachers = natsort.natsorted({teacher for *_, cell in lessons for teacher in cell.get("teachers", [])})
    nobody = translate("web.classes.no_teacher")

    if any(not cell.get("teachers") for *_, cell in lessons):
        teachers.append(nobody)

    book = Workbook()
    page = book.active
    page.title = translate("menu.main.tab.export.schedule")

    rows, separators = weekSheet(page, settings, translate)
    groups = groupsByName(settings)

    for col, teacher in enumerate(teachers, start=4):
        page[f"{letter(col)}1"] = teacher

    # Уже записанные уроки: (курс-источник, день, урок, преподаватель) — копия того же урока пропускается
    written = set()

    for course, day, lesson, cell in lessons:
        if (day, lesson) not in rows:
            continue

        root = copyRoot(copies, course)

        for teacher in cell.get("teachers", []) or [nobody]:
            if (root, day, lesson, teacher) in written:
                continue

            written.add((root, day, lesson, teacher))
            target = page.cell(row=rows[(day, lesson)], column=teachers.index(teacher) + 4)
            content = withNote(teacherColumnLabel(groups.get(root, {}), root), shared.get(root))
            target.value = f"{target.value}\n{content}" if target.value else content
            target.fill = solidFill(color(cell["subject"]))

    finish(page, separators)

    return book


def parseDate(text):
    """Дата из строки "ГГГГ-ММ-ДД"; строку нельзя разобрать — CourseDatesError."""
    try:
        return datetime.date.fromisoformat(text)

    except (TypeError, ValueError):
        raise CourseDatesError(text)


def generateCalendarEvents(answer, course_groups, day_grid, default_start, default_end):
    """Разворачивает недельное расписание в список уроков по конкретным датам.

    Параметры:
    - answer — расписание {курс: неделя} (как в answer.json);
    - course_groups — список курсов (custom_groups), из них берутся даты начала и конца;
    - day_grid — время уроков: day_grid[день недели][номер урока] -> строка времени;
    - default_start / default_end — даты ("ГГГГ-ММ-ДД") для курсов без своих дат.

    Для каждого дня между датами курса (включительно) берётся урок того же дня недели. Уроков
    в выходные отдельно не убирать: если в сетке у дня нет уроков, их нет и в неделе курса.
    Курсы, у которых конец раньше начала, пропускаются; неверная дата — CourseDatesError.
    Возвращает список кортежей (дата, номер урока, курс, предмет, «учителя через запятую»,
    время), отсортированный по дате, номеру урока и курсу.
    """
    groups = {group["name"]: group for group in course_groups}
    default_start = parseDate(default_start)
    default_end = parseDate(default_end)
    events = []

    for course, week in answer.items():
        group = groups.get(course, {})
        start = parseDate(group["start_date"]) if group.get("start_date") else default_start
        end = parseDate(group["end_date"]) if group.get("end_date") else default_end

        for offset in range(max(0, (end - start).days + 1)):
            current = start + datetime.timedelta(days=offset)
            weekday = current.weekday()

            # Этого дня нет в недельном расписании курса (неделя короче семи дней)
            if weekday >= len(week):
                continue

            times = day_grid[weekday] if weekday < len(day_grid) else []

            for lesson, cell in enumerate(week[weekday]):
                if isLesson(cell):
                    time = times[lesson] if lesson < len(times) else ""
                    events.append((current, lesson, course, cell.get("subject", ""), ", ".join(cell.get("teachers", [])), time))

    events.sort(key=lambda event: (event[0], event[1], event[2]))

    return events


def calendarWorkbook(settings, answer, translate):
    """Каждый урок по датам — между стартом и окончанием его курса. При неверных датах бросает CourseDatesError.

    Столбцы: дата, день недели, время, номер урока, курс, предмет, преподаватель;
    включён автофильтр, шапка закреплена. У уроков общих курсов к предмету дописана пометка
    (``jointNotes``): «Математика (вместе с Потоком 1)». Даты у копии свои — уроки идут с начала её потока.
    """
    events = generateCalendarEvents(
        answer, courseGroups(settings), dayGrid(settings),
        settings.get("calendar_start_date", CALENDAR_START), settings.get("calendar_end_date", CALENDAR_END)
    )
    notes = jointNotes(settings, translate)

    book = Workbook()
    page = book.active
    page.title = sheetName(translate("menu.main.tab.export.calendar"))
    page.append([translate(f"menu.main.tab.export.{key}") for key in ("date", "day", "time", "lesson", "course", "subject", "teacher")])

    for current, lesson, course, subject, teachers, time in events:
        page.append([current, translate(f"day.{current.weekday()}"), time, lesson + 1, course,
                     withNote(subject, notes.get(course)), teachers or translate("web.classes.no_teacher")])

    page.auto_filter.ref = page.dimensions
    finish(page, [])

    for row in page.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "DD.MM.YYYY"

    # finish() закрепляет «D2» (под недельные листы) — здесь нужна только строка заголовков
    page.freeze_panes = "A2"

    return book


# ---------------------------------------------------------------- варианты этапа

def variantsWorkbook(settings, ranked, columns, translate, title):
    """Варианты этапа с «Предпросмотра»: лист сравнения и по листу на каждый вариант.

    ``ranked`` — варианты от лучшего к худшему, как их отдаёт сервер (number, metrics,
    answer, issues, accepted, rejected, best, tied — см. variants.rankVariants, teachers,
    teacherChanges, staleStarted и movedStarted — см. ranking.rankedVariants); ``columns`` — [(ключ правила, подпись)] строк
    сравнения; ``title`` — название этапа для заголовка.

    Лист «Сравнение»: правила по строкам, варианты по столбцам (как таблица на странице),
    с пометками «лучший», «одинаково», «принят», «отклонён» (как на странице) и «нельзя принять»
    (variants.unacceptable: вариант собран до начала курсов, которые уже идут, или сдвигает урок
    идущего курса), под правилами — блок
    «Кто ведёт» (variants.teacherCourses). Лист варианта: дни по столбцам, время
    по строкам, в ячейке — уроки этого времени («ЕГЭ основной, Химия — Иванова»)
    и под каждым — что с ним не так (те же подписи, что на карточках недели).
    """
    book = Workbook()
    compareSheet(book.active, ranked, columns, translate, title, teacherCourses(settings, ranked))

    for item in ranked:
        variantSheet(book, settings, item, translate)

    return book


def compareSheet(page, ranked, columns, translate, title, courses=()):
    """Лист сравнения вариантов: шапка — варианты с пометками, дальше строка на правило.

    Если есть курсы ``courses`` (variants.teacherCourses), под правилами — заголовок «Кто ведёт»
    (web.preview.teachers_head) и строка на курс: в ячейке варианта — кто ведёт курс. Если вариант
    сменит преподавателя принятого курса, ячейка залита и в ней ещё «Сейчас ведёт: …»
    (web.preview.teacher_now), как подсказка на странице; столбцы вариантов тогда шире (TEACHER_WIDTH).
    """
    bold = Font(bold=True)
    grey = solidFill("EEEEEE")
    # Последняя строка листа: правила, затем заголовок блока «Кто ведёт» и строки его курсов
    last = len(columns) + 1 + (len(courses) + 1 if courses else 0)
    page.title = translate("web.export_variants.compare")
    page["A1"] = title
    page["A1"].font = bold

    for col, item in enumerate(ranked, start=2):
        # «Лучший» — только если у него нет «ничьей» с другими вариантами; при ничьей у каждого
        # из равных вариантов вместо «лучший» пишется «одинаково», как на странице
        marks = [translate("menu.main.tab.preview.best")] if item.get("best") and not item.get("tied") else []
        marks += [translate("menu.main.tab.preview.equal")] if item.get("tied") else []
        marks += [translate("menu.main.tab.preview.accepted")] if item.get("accepted") else []
        # Совпадает с принятым расписанием (ranking.rankedVariants), как пометка на странице
        marks += [translate("web.preview.same")] if item.get("same") else []
        marks += [translate("web.preview.rejected")] if item.get("rejected") else []
        # Принятие откажет (stages.staleStarted, stages.movedStarted)
        marks += [translate("web.export_variants.stale")] if unacceptable(item) else []
        page.cell(row=1, column=col, value="\n".join([f"{translate('menu.main.tab.run.variant')} {item['number']}"] + marks)).font = bold

        if item.get("rejected"):
            for row in range(1, last + 1):
                page.cell(row=row, column=col).fill = grey

    for row, (key, label) in enumerate(columns, start=2):
        page.cell(row=row, column=1, value=label).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for col, item in enumerate(ranked, start=2):
            penalty = customId(key)
            value = item["metrics"]["custom"].get(penalty, 0) if penalty is not None else item["metrics"][key]
            page.cell(row=row, column=col, value=value)

    if courses:
        teacherRows(page, ranked, translate, courses, len(columns) + 2)

    page.row_dimensions[1].height = 48
    page.column_dimensions["A"].width = 46
    page.freeze_panes = "B2"

    for col in range(2, len(ranked) + 2):
        page.column_dimensions[letter(col)].width = TEACHER_WIDTH if courses else VARIANT_WIDTH

        for row in range(1, last + 1):
            page.cell(row=row, column=col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    page.cell(row=2, column=1).font = bold


def teacherRows(page, ranked, translate, courses, first):
    """Блок «Кто ведёт» листа сравнения со строки ``first``: заголовок и строка на каждый курс ``courses``.

    Если вариант снимает преподавателя принятого курса (его удалили или запретили ему курс после
    сборки), в ячейке «без преподавателя» (web.classes.no_teacher), а не пусто.
    """
    page.cell(row=first, column=1, value=translate("web.preview.teachers_head")).font = Font(bold=True)

    for row, course in enumerate(courses, start=first + 1):
        page.cell(row=row, column=1, value=course).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for col, item in enumerate(ranked, start=2):
            before = [change["before"] for change in item.get("teacherChanges", []) if change["course"] == course]
            names = ", ".join(item.get("teachers", {}).get(course, []))
            target = page.cell(row=row, column=col, value=names or (translate("web.classes.no_teacher") if before else ""))

            # Вариант сменит преподавателя принятого курса: кто ведёт сейчас, как подсказка на странице
            if before:
                now = ", ".join(dict.fromkeys(name for names in before for name in names))
                target.value += "\n" + fill(translate("web.preview.teacher_now"), teacher=now)
                target.fill = solidFill(CHANGED_FILL)


def variantSheet(book, settings, item, translate):
    """Лист одного варианта: неделя как на странице — дни по столбцам, время уроков по строкам.

    Имя листа — «Вариант 2», у отклонённого — «Вариант 1 (отклонён)». Цвет ячейки — как цвет
    подписей на карточках недели: есть хоть одна проблема (подпись без "level": "warn") — вся
    ячейка красная; только некритичные подписи (например, «В тот же день: …») — жёлтая
    (тёмно-жёлтый, как --warn в app.css, чтобы читался на белом).
    """
    bold = Font(bold=True)
    grid, days = gridRows(settings)
    name = f"{translate('menu.main.tab.run.variant')} {item['number']}" + (f" ({translate('web.preview.rejected')})" if item.get("rejected") else "")
    sheet = book.create_sheet(sheetName(name))
    sheet["A1"] = translate("menu.main.tab.export.time")

    for col, day in enumerate(days, start=2):
        sheet.cell(row=1, column=col, value=translate(f"abbreviate.day.{day}")).font = bold

    # Время по порядку начала урока («9:00» раньше «18:00», хотя как текст наоборот)
    times = sorted({time for day in days for time in grid[day]}, key=lambda time: [int(part) for part in re.findall(r"\d+", time)])
    cells = variantCells(settings, item, translate, grid, days)

    for row, time in enumerate(times, start=2):
        sheet.cell(row=row, column=1, value=time).font = bold

        for col, day in enumerate(days, start=2):
            lessons = sorted(cells.get((day, time), []))

            if not lessons:
                continue

            text = "\n\n".join(text + "".join(f"\n  — {problemText(problem)}" for problem in problems) for _, _, text, problems in lessons)
            target = sheet.cell(row=row, column=col, value=text)
            target.alignment = Alignment(vertical="top", wrap_text=True)

            found = [problem for *_, problems in lessons for problem in problems]

            if any(not isWarning(problem) for problem in found):
                target.font = Font(color=PROBLEM_COLOR)

            elif found:
                target.font = Font(color=WARNING_COLOR)

    sheet.column_dimensions["A"].width = 13
    sheet.freeze_panes = "B2"

    for col in range(2, len(days) + 2):
        sheet.column_dimensions[letter(col)].width = 42

    for row in range(2, len(times) + 2):
        sheet.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="top")


def variantCells(settings, item, translate, grid, days):
    """Уроки варианта по ячейкам листа: {(день, время): [(линейка, предмет, подпись, проблемы)]}.

    Подпись — «ЕГЭ основной, Химия — Иванова», у общих курсов с пометкой (``jointNotes``):
    «… (вместе с Потоком 1)»; проблемы — подписи урока из variants.lessonIssues.
    Уроки вне текущей сетки пропускаются.
    """
    groups = groupsByName(settings)
    notes = jointNotes(settings, translate)
    cells = {}

    for course in natsort.natsorted(item["answer"]):
        group = groups.get(course, {})
        # Линейка как на вкладке «Курсы» (courses.lineOf): у курса потока — программа («ЕГЭ основной»),
        # у доп. курса — своя линейка («Семинар ОГЭ»), а не общая программа «Семинары» — иначе
        # два семинара по одному предмету на листе не различить
        line = (lineOf(group) or "") if group else ""

        for day, lesson, cell in lessonEntries(item["answer"], course):
            if day in days and lesson < len(grid[day]):
                problems = item.get("issues", {}).get(course, {}).get(f"{day}-{lesson}", [])
                text = withNote(lessonText(cell, translate, f"{line}, "), notes.get(course))
                cells.setdefault((day, grid[day][lesson]), []).append((line, cell["subject"], text, problems))

    return cells


def problemText(problem):
    """Текст подписи о проблеме урока: строка или {"text", "level"} (variants.lessonIssues)."""
    return problem["text"] if isinstance(problem, dict) else problem


def isWarning(problem):
    """Подпись некритичная (жёлтая на странице): {"text", "level": "warn"}. Строка и подпись
    без "level" — проблема (красная)."""
    return isinstance(problem, dict) and problem.get("level") == "warn"


# ---------------------------------------------------------------- список преподавателей

def teacherListWorkbook(settings, answer, translate, stages):
    """Преподаватели проекта: что они ведут и когда им удобно.

    Лист «Преподаватели» — строка на преподавателя (``teacherSheet``). Дальше по листу на каждый
    этап из ``stages`` ([(ключ, название)]) с отметками времени (``availabilitySheet``).
    Расписание (``answer``) может быть пустым — тогда уроков в неделю 0.
    """
    teachers = natsort.natsorted(settings.get("teachers", {}))
    book = Workbook()
    teacherSheet(book.active, settings, answer, translate, teachers)

    for stage, label in stages:
        sheet = book.create_sheet(sheetName(f"{translate('menu.main.tab.export.teacher_list_time')} — {label}"))
        availabilitySheet(sheet, settings, translate, teachers, stage)

    return book


def teacherSheet(page, settings, answer, translate, teachers):
    """Лист «Преподаватели»: предметы, курсы «ведёт», «может вести», «запрещено», сколько курсов
    ведёт и сколько уроков в неделю у него в принятом расписании.

    Общий курс источника и копий считается один раз: копии в списках нет (её преподавателя задаёт
    источник), у источника пометка «вместе с Потоком 2»; общий урок — тоже один урок в неделю.
    """
    heads = ["teacher", "subjects", "assigned", "may", "forbidden", "courses_count", "lessons_count"]
    page.title = translate("menu.main.tab.export.teacher_list_sheet")
    copies = jointCopies(settings)
    notes = jointNotes(settings, translate)

    # Уроки в неделю у каждого преподавателя в принятом расписании: {(курс-источник, день, урок)} —
    # урок копии совпадает с уроком источника и второй раз не считается
    lessons = {}

    for course, day, lesson, cell in allLessons(answer):
        for teacher in cell.get("teachers", []):
            lessons.setdefault(teacher, set()).add((copyRoot(copies, course), day, lesson))

    for col, key in enumerate(heads, start=1):
        page.cell(row=1, column=col, value=translate(f"menu.main.tab.export.teacher_list.{key}")).font = Font(bold=True)

    for row, name in enumerate(teachers, start=2):
        subjects = [item["subject"] for item in settings["teachers"][name].get("subjects", [])]
        states = teacherCourseStates(settings, name, subjects)
        lists = ["\n".join(withNote(course, notes.get(course)) for course in states[key]) for key in ("assigned", "may", "forbidden")]
        values = [name, ", ".join(subjects), *lists, len(states["assigned"]), len(lessons.get(name, ()))]

        for col, value in enumerate(values, start=1):
            page.cell(row=row, column=col, value=value).alignment = Alignment(vertical="top", wrap_text=True)

    for col, width in enumerate([26, 22, 48, 48, 40, 14, 14], start=1):
        page.column_dimensions[letter(col)].width = width

    page.freeze_panes = "B2"
    page.auto_filter.ref = f"A1:{letter(len(heads))}{len(teachers) + 1}"


def teacherCourseStates(settings, name, subjects):
    """{"assigned", "may", "forbidden": [курсы]} — как преподаватель относится к курсам своих предметов
    (как на странице «Преподаватели»). Курсов-копий в списках нет: их ведёт тот же, кто ведёт
    источник, и отдельным курсом преподавателю они не считаются.
    """
    states = {"assigned": [], "may": [], "forbidden": []}
    copies = jointCopies(settings)

    for group in courseGroups(settings):
        subject = courseSubject(group)

        if subject in subjects and group["name"] not in copies:
            state = teacherCourseState(settings, name, subject, group["name"])

            if state in states:
                states[state].append(group["name"])

    return states


def availabilitySheet(sheet, settings, translate, teachers, stage):
    """Лист отметок времени этапа: неделя по строкам (как в «Расписании преподавателей»),
    преподаватели по столбцам, в ячейке — «удобно», «может» или «не может» с цветом.
    """
    rows, separators = weekSheet(sheet, settings, translate)

    for col, name in enumerate(teachers, start=4):
        sheet.cell(row=1, column=col, value=name).font = Font(bold=True)
        cannot = {tuple(slot) for slot in cannotSlots(settings, name, stage)}
        possible = {tuple(slot) for slot in possibleSlots(settings, name, stage)}

        for slot, row in rows.items():
            kind = "cannot" if slot in cannot else "may" if slot in possible else "ok"
            text, rgb = MARKS[kind]
            target = sheet.cell(row=row, column=col, value=translate(text))
            target.fill = solidFill(rgb)

    finish(sheet, separators)
