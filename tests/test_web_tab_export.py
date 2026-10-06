"""Выгрузки вкладки «Экспорт» (`src/web/tabs/export.py`): Excel по курсам, по преподавателям,
календарь, варианты этапа, список преподавателей и весь проект (.zip).

Проверяется, что выгрузки — настоящие книги Excel с каждым уроком расписания, имена файлов,
понятные ошибки (нет расписания, нет вариантов, неверные даты курса, неизвестный вид).

Выгрузки скачиваются, как со страницы, запросом GET /api/project/<имя>/export/<вид> (а не через
POST …/action). Основа — `RealProjectCase` из `tests/real_project.py`: копия реального проекта
2026/27, «сегодня» 04.10.2026. Отметка «линейка присоединяется к Потоку N» в архиве проекта
(``test_project_archive_keeps_joint_marks``) — на проекте builders.jointProject.
"""

import tests  # первым: папка данных тестов вместо %APPDATA% (см. tests/__init__.py)
tests.requireIsolation()

import io
import json
import urllib.parse
import zipfile
from unittest import mock
from urllib.parse import unquote

import openpyxl

from src.modules.translate import translate
from src.modules.functions.stages import stageCourses
from src.modules.functions.variants import saveVariant
from src.web.core import app
from src.web.tabs import export as export_tab
from tests.builders import GENERIC, JOINT_LINE, jointCourse, jointProject
from tests.real_project import RealProjectCase


class ExportTabTests(RealProjectCase):
    """Вкладка «Экспорт»."""
    NAME = "__test_export__"

    def download(self, kind, **query):
        """Скачивает выгрузку `kind`; возвращает ответ."""
        return self.client.get(f"/api/project/{self.NAME}/export/{kind}", query_string=query)

    def test_project_archive(self):
        """«Экспорт» → «Весь проект»: ZIP с настройками, расписанием и версиями, который снова импортируется."""
        response = self.download("project")
        self.assertEqual(response.status_code, 200)
        self.assertIn(self.NAME, response.headers["Content-Disposition"])

        with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
            names = {name.replace("\\", "/") for name in archive.namelist()}

        self.assertTrue({"settings.json", "answer.json", "weights.json"} <= names)
        self.assertTrue(any(name.startswith("versions/") for name in names))

    def test_project_archive_keeps_joint_marks(self):
        """AC-36: «Весь проект» после setJoint («ЕГЭ основной» Потока 2 вместе с Потоком 1): в архиве
        у курсов-копий together_with = 1, а их уроки в answer.json равны урокам источников.
        """
        self.openProject(*jointProject())
        self.ok("setJoint", section=2, line=JOINT_LINE, source=1)

        with zipfile.ZipFile(io.BytesIO(self.download("project").data)) as archive:
            files = {name.replace("\\", "/"): name for name in archive.namelist()}
            settings = json.loads(archive.read(files["settings.json"]).decode("utf-8"))
            answer = json.loads(archive.read(files["answer.json"]).decode("utf-8"))

        groups = {group["name"]: group for group in settings["classes"]["custom_groups"]}

        for subject in ("Математика", "Русский язык"):
            copy_, source = jointCourse(2, subject), jointCourse(1, subject)
            self.assertEqual(groups[copy_].get("together_with"), 1, copy_)
            self.assertEqual(answer[copy_], answer[source], copy_)

    def test_excel_exports(self):
        """Выгрузки «по курсам», «по преподавателям» и «календарь» — настоящие книги Excel с данными
        расписания; без расписания — понятная ошибка; неизвестный вид — тоже.
        """
        for kind in ("courses", "teachers", "calendar"):
            response = self.download(kind)
            self.assertEqual(response.status_code, 200, kind)
            # Имя файла «<проект>-курсы.xlsx» и т. п.: русские буквы — в виде filename*=UTF-8''…
            filename = f"{self.NAME}-{translate(f'menu.main.tab.export.{kind}_file')}.xlsx"
            self.assertIn(urllib.parse.quote(filename), response.headers["Content-Disposition"], kind)
            book = openpyxl.load_workbook(io.BytesIO(response.data))
            text = " ".join(str(cell) for sheet in book.worksheets for row in sheet.iter_rows(values_only=True) for cell in row if cell)
            self.assertIn("Клепачева Юлия" if kind != "courses" else "Русский язык", text, kind)

        self.assertEqual(self.download("pdf").status_code, 400)

        self.save("answer.json", {})
        response = self.download("courses")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["error"], translate("web.common.no_schedule"))

    def test_excel_exports_contain_every_lesson(self):
        """В выгрузке по курсам — каждый урок расписания в столбце своего курса; по преподавателям —
        каждый урок у каждого его преподавателя (уроки без преподавателя — в своём столбце); в календаре —
        урок на каждую дату между началом и концом курса в его день недели.
        """
        import datetime

        answer = self.load("answer.json")
        settings = self.load("settings.json")
        cells = [(name, day, lesson, cell) for name, week in answer.items() for day, row in enumerate(week)
                 for lesson, cell in enumerate(row) if cell.get("subject", "#") != "#"]

        sheet = openpyxl.load_workbook(io.BytesIO(self.download("courses").data)).worksheets[0]
        header = [cell.value for cell in sheet[1]]
        filled = sum(1 for row in sheet.iter_rows(min_row=2, min_col=4, values_only=True) for value in row if value)
        self.assertEqual(filled, len(cells))
        self.assertTrue({name for name, _, _, _ in cells} <= set(header))

        sheet = openpyxl.load_workbook(io.BytesIO(self.download("teachers").data)).worksheets[0]
        lines = sum(len(str(value).split("\n")) for row in sheet.iter_rows(min_row=2, min_col=4, values_only=True) for value in row if value)
        self.assertEqual(lines, sum(max(1, len(cell.get("teachers", []))) for _, _, _, cell in cells))

        rows = list(openpyxl.load_workbook(io.BytesIO(self.download("calendar").data)).worksheets[0].iter_rows(min_row=2, values_only=True))
        course = "Поток 1 — ЕГЭ основной — Русский язык"
        group = next(item for item in settings["classes"]["custom_groups"] if item["name"] == course)
        start = datetime.date.fromisoformat(group["start_date"])
        end = datetime.date.fromisoformat(group.get("end_date") or settings["calendar_end_date"])
        weekdays = {day for name, day, _, _ in cells if name == course}
        expected = sum(1 for offset in range((end - start).days + 1) if (start + datetime.timedelta(days=offset)).weekday() in weekdays)
        mine = [row for row in rows if row[4] == course]
        self.assertEqual(len(mine), expected)
        self.assertTrue(all(start <= row[0].date() <= end for row in mine))

    def test_export_variants_to_excel(self):
        """«Скачать варианты» (книга Excel): лист сравнения (правила × варианты, с пометкой «отклонён»)
        и по листу недели на каждый вариант, с уроками и подписями о проблемах.
        """
        import io
        import openpyxl

        answer = self.load("answer.json")
        week = {name: answer[name] for name in stageCourses(self.load("settings.json"), "1") if name in answer}
        saveVariant(self.folder, "1", 1, week)
        saveVariant(self.folder, "1", 2, week)
        self.act("rejectVariant", stage="1", number=2)

        response = self.client.get(f"/api/project/{self.NAME}/export/variants?stage=1")
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.data))

        self.assertEqual(book.sheetnames, ["Сравнение", "Вариант 1", "Вариант 2 (отклонён)"])
        compare = book["Сравнение"]
        self.assertIn("отклонён", compare.cell(row=1, column=3).value)
        # Вариант 1 не с кем сравнивать (второй отклонён) — он «лучший»
        self.assertIn("лучший", compare.cell(row=1, column=2).value)
        self.assertEqual(compare.cell(row=2, column=1).value, "Неудобства, всего (меньше — лучше)")

        lessons = [cell.value for row in book["Вариант 1"].iter_rows(min_row=2, min_col=2) for cell in row if cell.value]
        self.assertTrue(any("ЕГЭ основной, Математика база — Файзиева Ангелина" in text for text in lessons))

        # Неизвестный поток — понятная ошибка, а не падение
        self.assertEqual(self.client.get(f"/api/project/{self.NAME}/export/variants?stage=9").status_code, 400)

    def test_export_teacher_list(self):
        """Выгрузка «Преподаватели»: у каждого — курсы «ведёт» / «может вести» и уроки в неделю,
        и по листу времени на каждый поток с отметками «удобно / может / не может».
        """
        import io
        import openpyxl

        response = self.client.get(f"/api/project/{self.NAME}/export/teacher_list")
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.data))

        self.assertEqual(book.sheetnames[0], "Преподаватели")
        self.assertTrue(all(name.startswith("Удобство — ") for name in book.sheetnames[1:]))
        rows = {row[0]: row for row in book["Преподаватели"].iter_rows(min_row=2, values_only=True)}
        self.assertIn("Поток 1 — ЕГЭ основной — Математика база", rows["Файзиева Ангелина"][2])
        self.assertGreater(rows["Файзиева Ангелина"][6], 0)
        marks = {cell for row in book[book.sheetnames[1]].iter_rows(min_row=2, min_col=4, values_only=True) for cell in row if cell}
        self.assertTrue(marks <= {"удобно", "может", "не может"} and "удобно" in marks)

    def test_export_variants_has_custom_rule_row(self):
        """В выгрузке вариантов есть строка своего правила с его значением, как на странице."""
        self.ok("savePenalty", penalty={"name": "Мутационное правило", "template": "same_day", "weight": 5, "params": {"first": "Химия", "second": "Биология"}})
        saveVariant(self.folder, "2", 1, self.stageVariant())
        body = self.variants()
        rule = body["penalties"][0]

        response = self.client.get(f"/api/project/{self.NAME}/export/variants", query_string={"stage": "2"})

        self.assertEqual(response.status_code, 200)
        compare = openpyxl.load_workbook(io.BytesIO(response.data))[translate("web.export_variants.compare")]
        labels = {compare.cell(row=row, column=1).value: row for row in range(2, compare.max_row + 1)}
        self.assertIn("Мутационное правило", labels)
        self.assertEqual(compare.cell(row=labels["Мутационное правило"], column=2).value, body["variants"][0]["metrics"]["custom"].get(rule["id"], 0))

    def test_export_file_names(self):
        """Имена скачиваемых книг: «<проект>-курсы.xlsx», «-преподаватели.xlsx», «-календарь.xlsx» (UTF-8)."""
        for kind, word in (("courses", "курсы"), ("teachers", "преподаватели"), ("calendar", "календарь")):
            response = self.client.get(f"/api/project/{self.NAME}/export/{kind}")
            self.assertEqual(response.status_code, 200, kind)
            header = response.headers["Content-Disposition"]
            self.assertIn("filename*=UTF-8''", header, kind)
            self.assertEqual(unquote(header.split("filename*=UTF-8''")[1].split(";")[0]), f"{self.NAME}-{word}.xlsx")
            response.close()

    def test_export_variants_without_variants(self):
        """Выгрузка вариантов этапа, у которого их ещё нет, — понятная ошибка «Вариантов пока нет — составьте
        их на шаге «Запуск»» (тот же текст, что на кнопке вкладки); неизвестный этап — общая ошибка: такой
        адрес страница не строит.
        """
        cases = (("2", translate("web.export_variants.none")), ("9", GENERIC))

        for stage, error in cases:
            response = self.client.get(f"/api/project/{self.NAME}/export/variants", query_string={"stage": stage})
            self.assertEqual((response.status_code, response.get_json()["error"]), (400, error), stage)

    def test_export_with_broken_course_dates(self):
        """Календарь при испорченной дате курса — понятное объяснение, а не падение."""
        settings = self.load("settings.json")
        settings["classes"]["custom_groups"][0]["start_date"] = "2026-13-45"
        self.save("settings.json", settings)

        response = self.client.get(f"/api/project/{self.NAME}/export/calendar")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.get_json()["error"], translate("menu.main.tab.export.invalid_course_dates"))

    def test_other_value_error_is_not_called_wrong_dates(self):
        """Ошибка сборки книги, не связанная с датами, не выдаётся за «неверные даты курсов»: это
        непредвиденная ошибка (500, общий текст, запись в журнал).
        """
        with mock.patch.dict(export_tab.SCHEDULE_BOOKS, {"courses": (mock.Mock(side_effect=ValueError("сломалось")), "x")}), \
                self.assertLogs(app.logger, "ERROR"):
            response = self.client.get(f"/api/project/{self.NAME}/export/courses")

        self.assertEqual((response.status_code, response.get_json()["error"]), (500, GENERIC))
