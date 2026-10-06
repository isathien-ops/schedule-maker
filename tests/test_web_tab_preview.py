"""Вкладка «Предпросмотр» (`src/web/tabs/preview.py`, `src/web/ranking.py`).

* ``VariantsListTests`` — ответ /api/project/<p>/variants: поля вариантов, отметка сборки,
  «принят» / «как сейчас», порядок (жёсткие нарушения — в конце), неизвестный этап;
* ``AcceptRejectTests`` — принятие варианта (вопросы о переставленных уроках, накладке
  и устаревших вариантах, удалённые преподаватели) и отклонение / возврат варианта;
* ``MovedLessonsTests`` — ``ranking.movedLessons``: какие уже принятые уроки переедут;
* ``JointAcceptTests`` — «линейка присоединяется к Потоку N» (.spec/joint-lines/SPEC.md): принятие
  варианта Потока 1 двигает копии и предупреждает о конфликтах в их потоке и о часах «не может»
  (AC-39), принятие варианта Потока 2 копии не трогает, копии, ждущие Поток 1, — причина joint_waiting;
* ``JointStaleVariantTests`` — вариант Потока 2, собранный до того, как Поток 1 сдвинули: копии
  в нём показываются и принимаются в нынешних часах источника, помехи названы в вопросе;
* ``TeacherChangeVariantsTests`` — «Смена преподавателя в подборе» (.spec/teacher-swap/SPEC.md):
  поля /variants ``teachers``, ``teacherChanges`` и ``teacherCourses`` (AC-16);
* ``TeacherChangeAcceptTests`` — принятие варианта, который меняет преподавателя курса: вопрос,
  отмена, принятие с force, копии (AC-18), накладка нового преподавателя у общего урока (AC-19),
  превышение лимита курсов (AC-20);
* ``TeacherChangeReviewTests`` — замечания проверки: у общего урока сменился только преподаватель,
  преподавателя из варианта удалили или запретили ему курс, абзац о смене — перед переставленными уроками;
* ``CopyBlockerAcceptTests`` — помеха в вопросе принятия — урок копии: составить заново советуют этап
  её источника (project.conflictStages);
* ``StaleStartedAcceptTests`` — вариант, собранный до начала курса, который уже идёт (stages.staleStarted):
  принятие отказывает даже с force, /variants отдаёт те же курсы в ``staleStarted``; курсу, который шёл
  уже при сборке (started.json), вариант сдвинул урок (stages.movedStarted) — тоже отказ, с уроком и причиной;
* ``StageClashAcceptTests`` — после подмены зафиксированного курса уроками из расписания у
  преподавателя два урока сразу: метрика teacherClash и вопрос web.preview.confirm_clash;
* ``LockedVariantTests`` — вариант, в котором зафиксированному курсу другой преподаватель: «Кто ведёт»
  и вопрос принятия (лимит курсов) видят преподавателя из расписания.

Основа — `RealProjectCase` из `tests/real_project.py`: копия реального проекта 2026/27, «сегодня»
04.10.2026, действия идут через сервер, как со страницы (POST /api/project/<имя>/action).
Проверяется, что сервер переспрашивает (ответ с "confirm") перед принятием варианта, которое
переставит уже принятые уроки или даст накладку, и сохраняет версию «Перед: принятие варианта».
"""

import tests  # первым: папка данных тестов вместо %APPDATA% (см. tests/__init__.py)
tests.requireIsolation()

import copy
import datetime
import io
import itertools
import json
import os
import unittest
from unittest import mock

import openpyxl

from src.modules.translate import tr, translate
from src.modules.functions.courses import setSectionEnd
from src.modules.functions.model import hasLessons
from src.modules.functions.stages import stageCourses
from src.modules.functions.variants import HARD_METRICS, METRICS, clearVariants, saveBuildStarted, saveVariant, variantsDir
from src.web import build, ranking
from src.web.project import slotText
from src.web.ranking import movedLessons
from tests.builders import (
    GENERIC, JOINT_LINE, LEVEL_LINE, OWN_LINE, courseWeek, jointCourse, jointProject, markCannot, markJoint, setHours, shiftStream, stageVariant
)
from tests.real_project import CHEM_OGE_2, RealProjectCase, keyPattern, lessons

EMPTY = {"subject": "#", "teachers": []}

# Курсы проекта builders.jointProject: источники Потока 1, копии и обычные курсы Потока 2
MATH_1, MATH_2 = jointCourse(1, "Математика"), jointCourse(2, "Математика")
RUS_1, RUS_2 = jointCourse(1, "Русский язык"), jointCourse(2, "Русский язык")
INF_2 = jointCourse(2, "Информатика")

# Уровень «ЕГЭ продвинутый» Математики Потока 1 и обычные курсы Математики Потока 2
MATH_LEVEL_1 = jointCourse(1, "Математика", LEVEL_LINE)
MATH_LEVEL_2, MATH_OWN_2 = jointCourse(2, "Математика", LEVEL_LINE), jointCourse(2, "Математика", OWN_LINE)

# «Смена преподавателя в подборе»: в принятом Потоке 1 «ЕГЭ основной» Математику ведёт OLD_TEACHER,
# в варианте — NEW_TEACHER (оба «могут вести» все курсы Математики проекта builders.jointProject)
OLD_TEACHER, NEW_TEACHER = "Математика #1", "Математика #3"

# Своя часть варианта Потока 2 (без копий): уроки курсов, которые не идут вместе с Потоком 1
OWN_2 = {
    INF_2: courseWeek("Информатика", "Информатика #1", (4, 1)),
    jointCourse(2, "Математика", LEVEL_LINE): courseWeek("Математика", "Математика #3", (1, 2), (3, 2)),
    jointCourse(2, "Русский язык", LEVEL_LINE): courseWeek("Русский язык", "Русский язык #3", (0, 2), (2, 2)),
    jointCourse(2, "Математика", OWN_LINE): courseWeek("Математика", "Математика #2", (4, 0)),
}


def week(slots, subject="Химия"):
    """Неделя 7 дней по 3 ячейки: уроки предмета в ячейках `slots`, остальные пустые."""
    return [[{"subject": subject, "teachers": []} if (day, lesson) in slots else dict(EMPTY) for lesson in range(3)] for day in range(7)]


def withTeacher(week, name):
    """Та же неделя курса (глубокая копия), но во всех её уроках преподаватель ``name``."""
    result = copy.deepcopy(week)

    for _, _, cell in lessons(result):
        cell["teachers"] = [name]

    return result


def paragraphWith(key, text):
    """Абзац вопроса ``text`` (абзацы разделены пустой строкой, как в acceptQuestion), в котором
    есть текст ``key`` из ru.hjson; нет такого абзаца — пустая строка.
    """
    return next((part for part in text.split("\n\n") if keyPattern(key).search(part)), "")


class VariantsListTests(RealProjectCase):
    """Список вариантов этапа для страницы."""
    NAME = "__test_variants__"

    def test_variants_answer_fields(self):
        """Ответ /variants содержит всё, на что опирается страница: hard, metrics, limit, build, penalties."""
        saveVariant(self.folder, "2", 1, self.stageVariant())
        self.ok("savePenalty", penalty={"name": "Химия и биология", "template": "same_day", "weight": 5, "params": {"first": "Химия", "second": "Биология"}})

        body = self.variants()

        self.assertEqual(body.get("hard"), list(HARD_METRICS))
        self.assertEqual(body["metrics"], [key for key, _ in METRICS])
        self.assertEqual(body["limit"], self.load("settings.json").get("max_courses_per_teacher", 5))
        self.assertTrue(body["build"])
        self.assertEqual([item["name"] for item in body["penalties"]], ["Химия и биология"])
        self.assertTrue(body["penalties"][0]["description"])
        self.assertEqual([item["number"] for item in body["variants"]], [1])
        self.assertIs(body["allTied"], False)

    def test_build_mark_ignores_accepted_and_rejected_files(self):
        """Принятие и отклонение не меняют «отметку сборки»: второй вариант принимается со старой отметкой."""
        saveVariant(self.folder, "2", 1, self.stageVariant())
        saveVariant(self.folder, "2", 2, self.stageVariant())
        build = self.variants()["build"]
        folder = variantsDir(self.folder, "2")
        later = datetime.datetime.now().timestamp() + 1000

        self.ok("rejectVariant", stage="2", number=2)
        os.utime(os.path.join(folder, "rejected.json"), (later, later))
        self.assertEqual(self.variants()["build"], build)

        self.ok("accept", stage="2", number=1, force=True, build=build)
        os.utime(os.path.join(folder, "accepted.json"), (later, later))
        self.assertEqual(self.variants()["build"], build)

        self.ok("accept", stage="2", number=2, force=True, build=build)
        self.assertEqual(self.load("stages/2.variants/accepted.json"), {"number": 2})

    def test_only_chosen_identical_variant_is_accepted(self):
        """Два одинаковых варианта, совпадающих с расписанием: «принят» только выбранный, другой — «как сейчас»."""
        saveVariant(self.folder, "2", 1, self.stageVariant())
        saveVariant(self.folder, "2", 2, self.stageVariant())

        flags = {item["number"]: (item["accepted"], item["same"]) for item in self.variants()["variants"]}
        self.assertEqual(flags, {1: (False, True), 2: (False, True)})

        with open(os.path.join(variantsDir(self.folder, "2"), "accepted.json"), "w", encoding="utf-8") as file:
            json.dump({"number": 1}, file)

        flags = {item["number"]: (item["accepted"], item["same"]) for item in self.variants()["variants"]}
        self.assertEqual(flags, {1: (True, False), 2: (False, True)})

    def test_order_puts_hard_violations_last(self):
        """Порядок вариантов: сначала без жёстких нарушений, затем по общему штрафу, затем по номеру —
        и на странице, и в листах выгрузки Excel.
        """
        for number in (1, 2, 3, 4):
            saveVariant(self.folder, "2", number, self.stageVariant())

        def metrics(missing, total):
            return dict({key: 0 for key, _ in METRICS}, missing=missing, total=total, custom={})

        # Вызовы идут по номерам вариантов: 1 — нехватка урока при малом штрафе, 2 и 3 — равные, 4 — лучший
        order = itertools.cycle([metrics(1, 0), metrics(0, 10), metrics(0, 10), metrics(0, 5)])

        with mock.patch.object(ranking, "variantMetrics", side_effect=lambda *args, **kwargs: copy.deepcopy(next(order))):
            items = self.variants()["variants"]
            response = self.client.get(f"/api/project/{self.NAME}/export/variants", query_string={"stage": "2"})

        self.assertEqual([item["number"] for item in items], [4, 2, 3, 1])
        # «Лучший» — вариант 4, ничьей с ним нет (страница и Excel берут эти поля как есть)
        self.assertEqual([(item["best"], item["tied"]) for item in items], [(True, False), (False, False), (False, False), (False, False)])
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.data))
        variant = translate("menu.main.tab.run.variant")
        self.assertEqual(book.sheetnames[1:], [f"{variant} {number}" for number in (4, 2, 3, 1)])
        compare = book[translate("web.export_variants.compare")]
        self.assertEqual([str(compare.cell(row=1, column=col).value).split("\n")[0] for col in range(2, 6)], [f"{variant} {number}" for number in (4, 2, 3, 1)])

    def test_variants_of_unknown_stage_are_empty(self):
        """Варианты без этапа или неизвестного этапа — пустой список, а не ошибка; «отметка сборки» пустая."""
        for query in ({}, {"stage": "9"}):
            data = self.client.get(f"/api/project/{self.NAME}/variants", query_string=query).get_json()
            self.assertEqual((data["variants"], data["build"]), ([], ""), query)


class AcceptRejectTests(RealProjectCase):
    """Принятие и отклонение варианта."""
    NAME = "__test_accept__"

    def clashVariant(self):
        """Вариант потока 2: всё как в расписании плюс новый курс, чей преподаватель в те же часы ведёт урок
        потока 1 (накладка с другим этапом).
        """
        answer = self.load("answer.json")
        variant = {name: answer[name] for name in stageCourses(self.load("settings.json"), "2") if name in answer}
        variant["Поток 2 — ЕГЭ основной — Русский язык"] = copy.deepcopy(answer["Поток 1 — ЕГЭ основной — Русский язык"])
        return variant

    def test_accept_asks_about_moved_lessons(self):
        """Вариант переносит принятый урок: в поле moved — 1, без force вопрос «переедут (1 урок)» с именем курса,
        расписание не меняется; с force урок встаёт на новое место.
        """
        variant = self.stageVariant()
        cell = variant[CHEM_OGE_2][3][1]
        variant[CHEM_OGE_2][3][1] = dict(EMPTY)
        variant[CHEM_OGE_2][2][2] = cell
        saveVariant(self.folder, "2", 1, variant)
        before = self.raw("answer.json")

        self.assertEqual(self.variants()["variants"][0]["moved"], 1)

        body = self.ok("accept", stage="2", number=1)
        moved = translate("web.preview.confirm_moved").replace("{lessons}", "1 урок").replace("{courses}", f"• {CHEM_OGE_2}")
        self.assertIn(moved, body["confirm"])
        self.assertTrue(body["confirm"].endswith(translate("web.preview.confirm_accept")))
        self.assertEqual(self.raw("answer.json"), before)

        self.ok("accept", stage="2", number=1, force=True)
        self.assertEqual([(day, lesson) for day, lesson, _ in lessons(self.load("answer.json")[CHEM_OGE_2])], [(2, 2)])

    def test_variant_that_only_adds_lesson_moves_nothing(self):
        """Вариант оставляет принятые уроки на месте и добавляет ещё один: moved = 0, принимается без вопроса."""
        variant = self.stageVariant()
        variant[CHEM_OGE_2][2][2] = {"subject": "Химия", "teachers": ["Баранникова Анна"]}
        saveVariant(self.folder, "2", 1, variant)

        self.assertEqual(self.variants()["variants"][0]["moved"], 0)
        body = self.ok("accept", stage="2", number=1)
        self.assertNotIn("confirm", body)

    def test_accept_drops_removed_teacher(self):
        """Преподаватель из варианта, которого уже нет в проекте, не попадает в расписание при принятии."""
        variant = self.stageVariant()
        variant[CHEM_OGE_2][3][1]["teachers"] = ["Баранникова Анна", "Удалённый"]
        saveVariant(self.folder, "2", 1, variant)

        self.ok("accept", stage="2", number=1, force=True)

        self.assertEqual(self.load("answer.json")[CHEM_OGE_2][3][1]["teachers"], ["Баранникова Анна"])

    def test_accept_with_clash_only_asks_about_clash(self):
        """Вариант ничего не переставляет, но ставит преподавателя на два урока сразу: вопрос только о накладке;
        с force вариант принимается, сохраняется версия и метка «принят».
        """
        saveVariant(self.folder, "2", 1, self.clashVariant())
        versions = self.versions()

        body = self.ok("accept", stage="2", number=1)
        self.assertIn(translate("web.preview.confirm_clash").split("{")[0], body["confirm"])
        self.assertNotIn(translate("web.preview.confirm_moved").split("{")[0], body["confirm"])
        self.assertNotIn("Поток 2 — ЕГЭ основной — Русский язык", self.load("answer.json"))

        self.ok("accept", stage="2", number=1, force=True)
        answer = self.load("answer.json")
        self.assertEqual(answer["Поток 2 — ЕГЭ основной — Русский язык"], self.clashVariant()["Поток 2 — ЕГЭ основной — Русский язык"])
        self.assertEqual(len(self.versions()), len(versions) + 1)
        self.assertTrue(self.state()["clashes"])

        data = self.client.get(f"/api/project/{self.NAME}/variants", query_string={"stage": "2"}).get_json()
        self.assertTrue(data["variants"][0]["accepted"])
        self.assertGreater(data["variants"][0]["metrics"]["teacherClash"], 0)

    def test_accept_refuses_unknown_number(self):
        """Принять несуществующий вариант или «номер» не числом нельзя; расписание не меняется."""
        saveVariant(self.folder, "2", 1, self.clashVariant())
        before = self.raw("answer.json")

        for number in (2, "abc", -1, None):
            self.assertEqual(self.refused("accept", stage="2", number=number, force=True), translate("web.error.no_variant"), number)

        self.assertEqual(self.raw("answer.json"), before)

    def test_reject_and_return_variant(self):
        """Отклонить вариант и вернуть его (номер и числом, и строкой); несуществующий вариант или этап — отказ."""
        saveVariant(self.folder, "2", 1, self.clashVariant())
        address = f"/api/project/{self.NAME}/variants"

        self.ok("rejectVariant", stage="2", number=1)
        self.assertTrue(self.client.get(address, query_string={"stage": "2"}).get_json()["variants"][0]["rejected"])
        self.ok("rejectVariant", stage="2", number="1", rejected=False)
        self.assertFalse(self.client.get(address, query_string={"stage": "2"}).get_json()["variants"][0]["rejected"])

        self.assertEqual(self.refused("rejectVariant", stage="2", number=5), translate("web.error.no_variant"))
        self.assertEqual(self.refused("rejectVariant", stage="9", number=1), GENERIC)

    def test_accept_asks_about_moved_lessons_and_stale_variants(self):
        """Принятие варианта, который двигает уже принятые уроки, требует подтверждения; устаревший
        вариант (от другой сборки) не принимается; после принятия сохраняется версия для отката.
        """
        stage = "2"
        settings = self.load("settings.json")
        answer = self.load("answer.json")
        moved = {}

        # Вариант, где один принятый урок перенесён на свободный день своего курса
        for name in stageCourses(settings, stage):
            if name in answer and lessons(answer[name]):
                week = json.loads(json.dumps(answer[name]))
                day, lesson, cell = lessons(week)[0]
                days = {d for d, _, _ in lessons(week)}
                free = next(d for d in range(5) if d not in days and len(week[d]) > lesson)
                week[free][lesson], week[day][lesson] = cell, {"subject": "#", "teachers": []}
                moved = {**{n: answer[n] for n in stageCourses(settings, stage) if n in answer}, name: week}
                break

        saveVariant(self.folder, stage, 1, moved)
        build = self.client.get(f"/api/project/{self.NAME}/variants?stage={stage}").get_json()["build"]

        self.assertIn("confirm", self.act("accept", stage=stage, number=1, build=build)[1])
        self.assertEqual(self.act("accept", stage=stage, number=1, build="old")[0], 400)
        self.assertEqual(self.act("accept", stage=stage, number=1, build=build, force=True)[0], 200)
        self.assertTrue(any(name.startswith("Перед: принятие варианта") for name in self.versions()))

    def test_reject_variant_marks_it_until_variants_are_rebuilt(self):
        """«Отклонить» помечает вариант (он остаётся в списке, но бледным), «Вернуть» снимает
        отметку, а новое составление вариантов начинается без отметок.
        """
        saveVariant(self.folder, "2", 1, {})
        saveVariant(self.folder, "2", 2, {})
        rejected = lambda: {item["number"]: item["rejected"] for item in self.client.get(f"/api/project/{self.NAME}/variants?stage=2").get_json()["variants"]}

        self.assertEqual(self.act("rejectVariant", stage="2", number=1)[0], 200)
        self.assertEqual(rejected(), {1: True, 2: False})
        self.act("rejectVariant", stage="2", number=1, rejected=False)
        self.assertEqual(rejected(), {1: False, 2: False})
        # Несуществующий вариант — ошибка
        self.assertEqual(self.act("rejectVariant", stage="2", number=7)[0], 400)

        self.act("rejectVariant", stage="2", number=2)
        clearVariants(self.folder, "2")
        saveVariant(self.folder, "2", 2, {})
        self.assertEqual(rejected(), {2: False})


class MovedLessonsTests(unittest.TestCase):
    """Какие уже принятые уроки переедут, если принять вариант."""
    def test_moved_lessons_are_removed_places(self):
        """Переставленные уроки — места из принятого расписания, которых нет в варианте; курс не из расписания пропускается."""
        answer = {"A": week([(0, 0), (1, 1)])}
        variant = {"A": week([(0, 0), (2, 2)]), "B": week([(0, 0)])}

        self.assertEqual(movedLessons(answer, variant, ["A", "B"]), [("A", (1, 1))])
        # Вариант только добавляет урок — ничего не переставлено
        self.assertEqual(movedLessons(answer, {"A": week([(0, 0), (1, 1), (2, 2)])}, ["A"]), [])

    def test_moved_lessons_follow_given_course_order(self):
        """Курсы перечисляются в переданном порядке (порядке курсов проекта), а не как попало:
        вопрос о переезде уроков при каждом принятии одинаковый.
        """
        answer = {"Б": week([(1, 1)]), "А": week([(0, 0)]), "В": week([(2, 2)])}
        variant = {"Б": week([(1, 2)]), "А": week([(0, 1)]), "В": week([(2, 0)])}

        self.assertEqual([course for course, _ in movedLessons(answer, variant, ["В", "А", "Б"])], ["В", "А", "Б"])


class JointAcceptTests(RealProjectCase):
    """Принятие вариантов, когда «ЕГЭ основной» Потока 2 идёт вместе с Потоком 1 (builders.jointProject)."""
    NAME = "__test_joint_accept__"

    def sourceVariant(self, conflict=False):
        """Поток 1 ещё не начался и принят, «ЕГЭ основной» Потока 2 идёт вместе с ним. Сохраняется
        вариант Потока 1, где «Математика» «ЕГЭ основной» стоит в другие часы (вт и чт, 1-й урок).

        ``conflict`` — в линейке копии в Потоке 2 во вторник 1-м уроком стоит «Информатика»,
        а «Математика» и «Информатика» — пара «нельзя одновременно». Возвращает вариант.
        """
        settings, answer = jointProject()

        shiftStream(settings, 1, "2026-10-12")

        if conflict:
            settings.setdefault("joint_subject_pairs", []).append(["Математика", "Информатика"])
            answer[INF_2] = courseWeek("Информатика", "Информатика #1", (1, 0))

        markJoint(settings, 2, JOINT_LINE, 1, answer)
        self.openProject(settings, answer)
        variant = {name: copy.deepcopy(answer[name]) for name in stageCourses(settings, "1") if name in answer}
        variant[MATH_1] = courseWeek("Математика", "Математика #1", (1, 0), (3, 0))
        saveVariant(self.folder, "1", 1, variant)

        return variant

    def test_accepting_source_variant_moves_copies(self):
        """AC-15: вариант Потока 1 двигает «ЕГЭ основной»: в вопросе — строка
        web.preview.confirm_joint_moved про Поток 2; после принятия уроки копии те же, что у источника.
        """
        variant = self.sourceVariant()

        body = self.ok("accept", stage="1", number=1)

        self.assertText("web.preview.confirm_joint_moved", body.get("confirm", ""))
        self.assertIn(JOINT_LINE, body["confirm"])

        self.ok("accept", stage="1", number=1, force=True)

        answer = self.load("answer.json")
        self.assertEqual(answer[MATH_1], variant[MATH_1])
        self.assertEqual(answer[MATH_2], answer[MATH_1])

    def test_accepting_source_variant_names_conflicts_of_copies(self):
        """AC-15: новые часы копии нарушают в Потоке 2 «нельзя одновременно» — в вопросе
        web.preview.confirm_joint_conflicts с курсом Потока 2; принять всё равно можно, копия
        встаёт в часы источника.
        """
        variant = self.sourceVariant(conflict=True)

        body = self.ok("accept", stage="1", number=1)

        self.assertText("web.preview.confirm_joint_conflicts", body.get("confirm", ""))
        self.assertIn(INF_2, body["confirm"])

        self.ok("accept", stage="1", number=1, force=True)
        self.assertEqual(self.load("answer.json")[MATH_2], variant[MATH_1])

    def test_accepting_source_variant_names_cannot_hours_of_copies(self):
        """AC-39: вариант Потока 1 сдвигает общую «Математику» на вт и чт 1-м уроком, а у «Математика #1»
        на этапе «Поток 2» во вт 1-м уроком «не может». В вопросе — абзац web.preview.confirm_joint_cannot
        с преподавателем и этим днём и часом. Ср 1-м уроком тоже «не может», но там общий урок уже стоит,
        и вариант его оттуда уводит, — не названа. Принять можно: копия встаёт в часы источника.
        """
        variant = self.sourceVariant()
        settings = self.load("settings.json")
        markCannot(settings, "Математика #1", "2", (1, 0), (2, 0))
        self.save("settings.json", settings)

        body = self.ok("accept", stage="1", number=1)

        confirm = body.get("confirm", "")
        self.assertIn(slotText(settings, 1, 0), confirm)
        self.assertIn("Математика #1", confirm)
        self.assertNotIn(slotText(settings, 2, 0), confirm)
        self.assertText("web.preview.confirm_joint_cannot", confirm)

        self.ok("accept", stage="1", number=1, force=True)
        self.assertEqual(self.load("answer.json")[MATH_2], variant[MATH_1])

    def test_accepting_copy_stream_keeps_copies(self):
        """AC-25: вариант Потока 2 без копий — принимается без вопроса «переставит уроки», копии
        в расписании остаются как у источника, вариант отмечен «принят».
        """
        settings, answer = jointProject()
        markJoint(settings, 2, JOINT_LINE, 1, answer)
        self.openProject(settings, answer)
        saveVariant(self.folder, "2", 1, OWN_2)

        body = self.ok("accept", stage="2", number=1)

        self.assertNotIn("confirm", body)
        answer = self.load("answer.json")
        self.assertEqual((answer[MATH_2], answer[RUS_2]), (answer[MATH_1], answer[RUS_1]))
        self.assertEqual({name: answer[name] for name in OWN_2}, OWN_2)
        self.assertTrue(self.variants("2")["variants"][0]["accepted"])

    def test_accepting_stale_copies_keeps_source_lessons(self):
        """AC-25: в варианте Потока 2 у копии другие уроки (старая сборка) — после принятия копия
        в расписании та же, что у источника.
        """
        settings, answer = jointProject()
        markJoint(settings, 2, JOINT_LINE, 1, answer)
        self.openProject(settings, answer)
        saveVariant(self.folder, "2", 1, dict(OWN_2, **{MATH_2: courseWeek("Математика", "Математика #1", (1, 0), (3, 0))}))

        self.ok("accept", stage="2", number=1, force=True)

        answer = self.load("answer.json")
        self.assertEqual(answer[MATH_2], answer[MATH_1])

    def test_waiting_copies_are_not_missing(self):
        """AC-13 (данные): Поток 1 не принят. В «Предпросмотре» Потока 2 у копий причина
        joint_waiting с номером Потока 1 (и никакой другой); Поток 2 можно составить и принять,
        копии в нём пустые.
        """
        settings, answer = jointProject(accepted=False)
        markJoint(settings, 2, JOINT_LINE, 1, answer)
        self.openProject(settings, answer)
        saveVariant(self.folder, "2", 1, {INF_2: OWN_2[INF_2]})

        problems = self.variants("2")["variants"][0]["problems"]

        for name in (MATH_2, RUS_2):
            reasons = [(item["reason"], item.get("number")) for item in problems if item["course"] == name]
            self.assertEqual(reasons, [("joint_waiting", 1)], name)

        with mock.patch.object(build, "SOLVER", os.path.abspath(__file__)), mock.patch.object(build, "startJob") as start:
            self.ok("run", stage="2")

        start.assert_called_once()
        self.ok("accept", stage="2", number=1, force=True)
        answer = self.load("answer.json")
        self.assertFalse(hasLessons(answer, MATH_2) or hasLessons(answer, RUS_2))
        self.assertEqual(answer[INF_2], OWN_2[INF_2])


class JointStaleVariantTests(RealProjectCase):
    """Вариант Потока 2 собран, когда общие уроки «ЕГЭ основной» стояли в другие часы, а потом Поток 1
    сдвинули (замечание проверки). Вариант показывается и принимается с копиями в нынешних часах
    источника, а помехи от новых часов названы в вопросе. Проект — ``builders.jointProject``.
    """
    NAME = "__test_joint_stale__"

    def staleVariant(self):
        """Поток 1 принят, «ЕГЭ основной» Потока 2 идёт вместе с ним, «Математика» и «Информатика» —
        пара «нельзя». Вариант Потока 2 сохранён со старыми копиями (Математика в пн и ср 1-м уроком)
        и «Информатикой» во вт 1-м уроком; после этого Математику Потока 1 сдвинули на вт и чт.
        """
        settings, answer = jointProject()
        settings.setdefault("joint_subject_pairs", []).append(["Математика", "Информатика"])
        markJoint(settings, 2, JOINT_LINE, 1, answer)
        variant = dict(copy.deepcopy({name: answer[name] for name in (MATH_2, RUS_2)}), **OWN_2)
        variant[INF_2] = courseWeek("Информатика", "Информатика #1", (1, 0))
        answer[MATH_1] = courseWeek("Математика", "Математика #1", (1, 0), (3, 0))
        markJoint(settings, 2, JOINT_LINE, 1, answer)
        self.openProject(settings, answer)
        saveVariant(self.folder, "2", 1, variant)

    def test_variant_shows_copies_as_now(self):
        """В /variants копия стоит в нынешних часах источника (без пунктира «новое»), у копий нет подписей
        уроков, а «уровни врозь» и пары посчитаны по этим часам.
        """
        self.staleVariant()
        answer = self.load("answer.json")

        item = self.variants("2")["variants"][0]

        self.assertEqual((item["answer"][MATH_2], item["answer"][RUS_2]), (answer[MATH_1], answer[RUS_1]))
        self.assertNotIn(MATH_2, item["issues"])
        self.assertEqual(item["moved"], 0)

    def test_accepting_stale_variant_names_conflicts(self):
        """Принятие без force — вопрос web.preview.confirm_joint_stale с «Информатикой» (она теперь
        одновременно с копией Математики); с force копия как у источника, уроки варианта приняты.
        """
        self.staleVariant()

        body = self.ok("accept", stage="2", number=1)

        self.assertText("web.preview.confirm_joint_stale", body.get("confirm", ""))
        self.assertIn(INF_2, body["confirm"])

        self.ok("accept", stage="2", number=1, force=True)

        answer = self.load("answer.json")
        self.assertEqual(answer[MATH_2], answer[MATH_1])
        self.assertEqual(answer[INF_2], courseWeek("Информатика", "Информатика #1", (1, 0)))

    def test_waiting_copies_in_variant_follow_source(self):
        """Вариант собран, когда Поток 1 ещё не был принят (копий в нём нет); теперь Поток 1 принят —
        в /variants копии с уроками источника, ждущих копий в problems нет.
        """
        settings, answer = jointProject()
        markJoint(settings, 2, JOINT_LINE, 1, answer)
        self.openProject(settings, answer)
        saveVariant(self.folder, "2", 1, OWN_2)

        item = self.variants("2")["variants"][0]

        self.assertEqual(item["answer"][MATH_2], answer[MATH_1])
        self.assertFalse([problem for problem in item["problems"] if problem["reason"] == "joint_waiting"])


class TeacherSwapCase(RealProjectCase):
    """База тестов «Смена преподавателя в подборе» (.spec/teacher-swap/SPEC.md) на проекте
    ``builders.jointProject``: Поток 1 принят и ещё не идёт (начинается 12.10.2026, «сегодня» 04.10),
    поэтому принятие варианта Потока 1 может менять все его курсы. «ЕГЭ основной» Математику Потока 1
    (``MATH_1``, пн и ср 1-м уроком) ведёт ``OLD_TEACHER``; ``NEW_TEACHER`` в эти часы свободен.
    Сам тестов не содержит.
    """


    def openSwapProject(self, joint=False, limit=None, extra=None):
        """Подменяет данные тестового проекта на ``jointProject`` с Потоком 1, который ещё не идёт, и открывает его.

        ``joint`` — «ЕГЭ основной» Потока 2 идёт вместе с Потоком 1 (``markJoint``: копии ``MATH_2``
        и ``RUS_2`` с уроками источников); ``limit`` — лимит курсов на преподавателя; ``extra`` —
        {курс: неделя}, которые добавить в принятое расписание. Возвращает (settings, answer).
        """
        settings, answer = jointProject()

        shiftStream(settings, 1, "2026-10-12")

        if limit is not None:
            settings["max_courses_per_teacher"] = limit

        answer.update(copy.deepcopy(extra or {}))

        if joint:
            markJoint(settings, 2, JOINT_LINE, 1, answer)

        self.openProject(settings, answer)

        return settings, answer

    def stageOne(self, **weeks):
        """Вариант Потока 1, совпадающий с принятым расписанием, где у курсов из ``weeks``
        ({курс: неделя}) свои недели. Имена курсов — ключи словаря ``weeks`` через ``**{…}``.
        """
        variant = self.stageVariant("1")
        variant.update(copy.deepcopy(weeks))

        return variant

    def swapVariant(self):
        """Вариант Потока 1, который меняет только преподавателя ``MATH_1``: те же часы, ``NEW_TEACHER``."""
        return self.stageOne(**{MATH_1: withTeacher(self.load("answer.json")[MATH_1], NEW_TEACHER)})

    def assertNotOverLimit(self, stage, *variants):
        """AC-20 (сторож): ``variants.teacherOverLimit`` пуст у каждого варианта ``variants`` этапа ``stage``."""
        from src.modules.functions.variants import teacherOverLimit

        settings, answer = self.load("settings.json"), self.load("answer.json")

        for variant in variants:
            self.assertEqual(list(teacherOverLimit(settings, answer, stage, variant)), [])


class TeacherChangeVariantsTests(TeacherSwapCase):
    """Поля /variants «кто ведёт» (SPEC «Сервер и страница», DATA_CONTRACT §7.6): у варианта
    ``teachers`` и ``teacherChanges``, в ответе ``teacherCourses``.
    """
    NAME = "__test_teacher_variants__"

    def test_variant_names_changed_teacher(self):
        """AC-16: вариант 1 меняет преподавателя ``MATH_1`` (``OLD_TEACHER`` → ``NEW_TEACHER``, часы те же),
        вариант 2 — как принятое расписание. У варианта 1 ``teacherChanges`` — ровно одна запись
        {course, subject, before, after}, у варианта 2 — пусто; ``teachers[MATH_1]`` у каждого свой;
        в ``teacherCourses`` только ``MATH_1``: у остальных курсов преподаватель везде один.
        Лимит курсов оба варианта не превышают.
        """
        self.openSwapProject()
        swap, same = self.swapVariant(), self.stageVariant("1")
        saveVariant(self.folder, "1", 1, swap)
        saveVariant(self.folder, "1", 2, same)

        body = self.variants("1")
        items = {item["number"]: item for item in body["variants"]}

        self.assertEqual(body.get("teacherCourses"), [MATH_1])
        self.assertEqual(items[1].get("teacherChanges"), [{"course": MATH_1, "subject": "Математика", "before": [OLD_TEACHER], "after": [NEW_TEACHER]}])
        self.assertEqual(items[2].get("teacherChanges"), [])
        self.assertEqual((items[1]["teachers"][MATH_1], items[2]["teachers"][MATH_1]), ([NEW_TEACHER], [OLD_TEACHER]))
        self.assertEqual(items[1]["teachers"][MATH_LEVEL_1], items[2]["teachers"][MATH_LEVEL_1])
        self.assertLessEqual(set(items[1]["teachers"]), set(stageCourses(self.load("settings.json"), "1")))
        self.assertNotOverLimit("1", swap, same)

    def test_accepted_variant_has_no_changes(self):
        """AC-16: после принятия варианта 1 (с force) у него ``teacherChanges == []``, а вариант 2 теперь
        вернул бы прежнего преподавателя: ``NEW_TEACHER`` → ``OLD_TEACHER``; ``MATH_1`` по-прежнему
        в ``teacherCourses``.
        """
        self.openSwapProject()
        saveVariant(self.folder, "1", 1, self.swapVariant())
        saveVariant(self.folder, "1", 2, self.stageVariant("1"))

        self.ok("accept", stage="1", number=1, force=True)
        body = self.variants("1")
        items = {item["number"]: item for item in body["variants"]}

        self.assertTrue(items[1]["accepted"])
        self.assertEqual(items[1].get("teacherChanges"), [])
        self.assertEqual(items[2].get("teacherChanges"), [{"course": MATH_1, "subject": "Математика", "before": [NEW_TEACHER], "after": [OLD_TEACHER]}])
        self.assertEqual(body.get("teacherCourses"), [MATH_1])

    def test_copies_and_new_courses_are_not_changes(self):
        """AC-16: вариант Потока 2, где «ЕГЭ основной» идёт вместе с Потоком 1. В сохранённом варианте
        у копии ``MATH_2`` другой преподаватель (старая сборка) — копия не в счёт: её нет ни
        в ``teacherChanges``, ни в ``teachers``, ни в ``teacherCourses``. «ЕГЭ продвинутый» Математику
        Потока 2 (её нет в расписании) варианты дают разным преподавателям: она в ``teacherCourses``,
        но не в ``teacherChanges`` — менять в расписании нечего.
        """
        self.openSwapProject(joint=True)
        first = dict(copy.deepcopy(OWN_2), **{MATH_2: courseWeek("Математика", "Математика #2", (0, 0), (2, 0))})
        second = dict(copy.deepcopy(first), **{MATH_LEVEL_2: withTeacher(OWN_2[MATH_LEVEL_2], OLD_TEACHER)})
        saveVariant(self.folder, "2", 1, first)
        saveVariant(self.folder, "2", 2, second)

        body = self.variants("2")
        items = {item["number"]: item for item in body["variants"]}

        self.assertEqual(body.get("teacherCourses"), [MATH_LEVEL_2])

        for number, teacher in ((1, NEW_TEACHER), (2, OLD_TEACHER)):
            self.assertEqual(items[number].get("teacherChanges"), [], number)
            self.assertEqual(items[number]["teachers"][MATH_LEVEL_2], [teacher], number)
            self.assertFalse({MATH_2, RUS_2} & set(items[number]["teachers"]), number)

    def test_rejected_variants_are_not_compared(self):
        """AC-16: ``teacherCourses`` сравнивает только не отклонённые варианты (и принятое расписание):
        вариант со сменой преподавателя отклонён, остался вариант как в расписании — список пуст;
        вариант вернули — ``MATH_1`` снова в нём.
        """
        self.openSwapProject()
        saveVariant(self.folder, "1", 1, self.swapVariant())
        saveVariant(self.folder, "1", 2, self.stageVariant("1"))

        self.assertEqual(self.variants("1").get("teacherCourses"), [MATH_1])
        self.ok("rejectVariant", stage="1", number=1)
        self.assertEqual(self.variants("1").get("teacherCourses"), [])
        self.ok("rejectVariant", stage="1", number=1, rejected=False)
        self.assertEqual(self.variants("1").get("teacherCourses"), [MATH_1])


class TeacherChangeAcceptTests(TeacherSwapCase):
    """Принятие варианта, который меняет преподавателя курса (SPEC «Принятие», Р-9)."""
    NAME = "__test_teacher_accept__"

    def test_accept_asks_about_new_teacher(self):
        """AC-18: вариант меняет только преподавателя ``MATH_1`` (часы те же). Без force — вопрос
        с абзацем web.preview.confirm_teacher и строкой «MATH_1: OLD_TEACHER → NEW_TEACHER»; это «Отмена»:
        answer.json и settings.json не меняются. С force — во всех уроках ``MATH_1`` ``NEW_TEACHER``,
        в settings «ведёт» и «может вести» как были, сохранена версия «Перед: принятие варианта».
        """
        self.openSwapProject()
        variant = self.swapVariant()
        saveVariant(self.folder, "1", 1, variant)
        answer, settings = self.raw("answer.json"), self.raw("settings.json")
        teachers = self.load("settings.json")["teachers"]

        body = self.ok("accept", stage="1", number=1)

        self.assertText("web.preview.confirm_teacher", body.get("confirm", ""))
        self.assertIn(f"{MATH_1}: {OLD_TEACHER} → {NEW_TEACHER}", body["confirm"])
        self.assertTrue(body["confirm"].endswith(translate("web.preview.confirm_accept")))
        self.assertEqual((self.raw("answer.json"), self.raw("settings.json")), (answer, settings))

        self.ok("accept", stage="1", number=1, force=True)

        self.assertEqual({tuple(cell["teachers"]) for _, _, cell in lessons(self.load("answer.json")[MATH_1])}, {(NEW_TEACHER,)})
        self.assertEqual(self.load("settings.json")["teachers"], teachers)
        self.assertTrue(any(name.startswith("Перед: принятие варианта") for name in self.versions()))
        self.assertNotOverLimit("1", variant)

    def test_accept_names_copies_of_changed_course(self):
        """AC-18: у ``MATH_1`` есть копия в Потоке 2 — в вопросе ещё и абзац web.preview.confirm_teacher_joint
        с номером потока-копии; после принятия копию ведёт ``NEW_TEACHER`` (``syncJointAnswer``).
        """
        self.openSwapProject(joint=True)
        saveVariant(self.folder, "1", 1, self.swapVariant())

        body = self.ok("accept", stage="1", number=1)

        self.assertText("web.preview.confirm_teacher", body.get("confirm", ""))
        self.assertText("web.preview.confirm_teacher_joint", body["confirm"])
        self.assertRegex(paragraphWith("web.preview.confirm_teacher_joint", body["confirm"]), r"(?<![#\d])2(?!\d)")

        self.ok("accept", stage="1", number=1, force=True)

        answer = self.load("answer.json")
        self.assertEqual({tuple(cell["teachers"]) for _, _, cell in lessons(answer[MATH_2])}, {(NEW_TEACHER,)})
        self.assertEqual(answer[MATH_2], answer[MATH_1])

    def test_variant_without_teacher_change_does_not_ask_about_teachers(self):
        """AC-18: вариант переносит урок ``MATH_1`` (пн → вт), преподаватель тот же — вопрос только
        о переставленных уроках, абзацев о смене преподавателя нет.
        """
        self.openSwapProject()
        saveVariant(self.folder, "1", 1, self.stageOne(**{MATH_1: courseWeek("Математика", OLD_TEACHER, (1, 0), (2, 0))}))

        confirm = self.ok("accept", stage="1", number=1).get("confirm", "")

        self.assertText("web.preview.confirm_moved", confirm)
        self.assertNoText("web.preview.confirm_teacher", confirm)
        self.assertNoText("web.preview.confirm_teacher_joint", confirm)
        self.assertNotIn("→", confirm)

    def test_new_teacher_busy_at_shared_lesson(self):
        """AC-19: ``MATH_1`` идёт вместе с Потоком 2, вариант оставляет его часы, но меняет преподавателя
        на ``NEW_TEACHER``, а тот в Потоке 2 в пн 1-м уроком ведёт «ОГЭ» Математику. ``joint.copyConflicts``
        находит у копии накладку (причина busy) — в вопросе абзац web.preview.confirm_joint_conflicts
        с этим курсом.
        """
        self.openSwapProject(joint=True, extra={MATH_OWN_2: courseWeek("Математика", NEW_TEACHER, (0, 0))})
        saveVariant(self.folder, "1", 1, self.swapVariant())

        confirm = self.ok("accept", stage="1", number=1).get("confirm", "")

        self.assertText("web.preview.confirm_joint_conflicts", confirm)
        self.assertIn(tr("menu.main.tab.classes.slot_busy", detail=MATH_OWN_2), paragraphWith("web.preview.confirm_joint_conflicts", confirm))

    def test_accept_names_teacher_over_limit(self):
        """AC-20: лимит 2 курса; ``NEW_TEACHER`` уже ведёт «ОГЭ» Математику Потока 2, а вариант Потока 1
        (файл правили руками) даёт ему ещё ``MATH_1`` и «ЕГЭ продвинутый» Математику — 3 курса.
        В вопросе принятия абзац web.preview.confirm_over_limit с его именем.
        """
        self.openSwapProject(limit=2, extra={MATH_OWN_2: courseWeek("Математика", NEW_TEACHER, (4, 2))})
        saveVariant(self.folder, "1", 1, self.stageOne(**{
            MATH_1: withTeacher(self.load("answer.json")[MATH_1], NEW_TEACHER),
            MATH_LEVEL_1: courseWeek("Математика", NEW_TEACHER, (1, 0), (3, 0)),
        }))

        confirm = self.ok("accept", stage="1", number=1).get("confirm", "")

        self.assertText("web.preview.confirm_over_limit", confirm)
        self.assertIn(NEW_TEACHER, paragraphWith("web.preview.confirm_over_limit", confirm))


class TeacherChangeReviewTests(TeacherSwapCase):
    """Замечания проверки смены преподавателя: у общего урока сменился только преподаватель (часы те же),
    преподавателя из варианта удалили или запретили ему курс после сборки, порядок абзацев вопроса.
    """
    NAME = "__test_teacher_review__"

    def test_stale_copy_with_new_teacher_only(self):
        """Вариант Потока 2 собран, пока общие уроки вёл ``OLD_TEACHER``; в нём «ОГЭ» Математику ведёт
        ``NEW_TEACHER`` в пн 1-м уроком. Потом Поток 1 приняли с ``NEW_TEACHER`` в те же часы. В вопросе
        принятия абзац web.preview.confirm_joint_stale с накладкой, и его текст говорит о смене
        преподавателя, а не только о часах.
        """
        _, answer = self.openSwapProject(joint=True)
        variant = dict(copy.deepcopy(OWN_2), **{MATH_2: answer[MATH_2], RUS_2: answer[RUS_2]})
        variant[MATH_OWN_2] = courseWeek("Математика", NEW_TEACHER, (0, 0))
        saveVariant(self.folder, "2", 1, variant)
        saveVariant(self.folder, "1", 1, self.swapVariant())
        self.ok("accept", stage="1", number=1, force=True)

        confirm = self.ok("accept", stage="2", number=1).get("confirm", "")

        self.assertText("web.preview.confirm_joint_stale", confirm)
        self.assertIn(tr("menu.main.tab.classes.slot_busy", detail=MATH_OWN_2), paragraphWith("web.preview.confirm_joint_stale", confirm))
        self.assertIn("преподавател", translate("web.preview.confirm_joint_stale").split("\n")[0])

    def test_cannot_hours_of_copy_with_new_teacher_only(self):
        """Вариант Потока 1 оставляет часы ``MATH_1`` и меняет преподавателя на ``NEW_TEACHER``, а у него
        в Потоке 2 на эти часы «не может». В вопросе абзац web.preview.confirm_joint_cannot с ним; уроки
        не сдвигаются (абзаца web.preview.confirm_joint_moved нет), и текст не говорит, что они сдвинутся.
        """
        settings, _ = self.openSwapProject(joint=True)
        markCannot(settings, NEW_TEACHER, "2", (0, 0), (2, 0))
        self.save("settings.json", settings)
        saveVariant(self.folder, "1", 1, self.swapVariant())

        confirm = self.ok("accept", stage="1", number=1).get("confirm", "")

        self.assertIn(NEW_TEACHER, paragraphWith("web.preview.confirm_joint_cannot", confirm))
        self.assertNoText("web.preview.confirm_joint_moved", confirm)
        self.assertNotIn("сдвин", translate("web.preview.confirm_joint_cannot"))

    def test_deleted_new_teacher_is_a_change(self):
        """Вариант даёт ``MATH_1`` ``NEW_TEACHER``, а его удалили: уроки курса в варианте остались без
        преподавателя. Это смена «``OLD_TEACHER`` → без преподавателя»: она в ``teacherChanges``
        и ``teacherCourses``, принятие без force спрашивает (web.preview.confirm_teacher).
        """
        self.openSwapProject()
        saveVariant(self.folder, "1", 1, self.swapVariant())
        self.ok("deleteTeacher", name=NEW_TEACHER, force=True)

        body = self.variants("1")

        self.assertEqual(body["variants"][0]["teacherChanges"], [{"course": MATH_1, "subject": "Математика", "before": [OLD_TEACHER], "after": []}])
        self.assertEqual(body["teacherCourses"], [MATH_1])

        confirm = self.ok("accept", stage="1", number=1).get("confirm", "")

        self.assertText("web.preview.confirm_teacher", confirm)
        self.assertIn(f"{MATH_1}: {OLD_TEACHER} → {translate('web.classes.no_teacher')}", confirm)

    def test_forbidden_new_teacher_does_not_come_back(self):
        """Вариант даёт ``MATH_1`` ``NEW_TEACHER``, а потом ему этот курс запретили (карточка преподавателя).
        Из варианта он снят (вариант 2, где курс ведёт ``OLD_TEACHER``, не меняется), принятие спрашивает
        о смене; после принятия с force он курс не ведёт.
        """
        self.openSwapProject()
        saveVariant(self.folder, "1", 1, self.swapVariant())
        saveVariant(self.folder, "1", 2, self.stageVariant("1"))
        same = self.raw(os.path.join(variantsDir(self.folder, "1"), "2.json"))

        while MATH_1 not in self.load("settings.json")["teachers"][NEW_TEACHER].get("forbidden", []):
            self.ok("cycleCourse", name=NEW_TEACHER, subject="Математика", course=MATH_1)

        self.assertNotIn(NEW_TEACHER, self.variants("1")["variants"][0]["teachers"].get(MATH_1, []))
        self.assertEqual(self.raw(os.path.join(variantsDir(self.folder, "1"), "2.json")), same)
        self.assertText("web.preview.confirm_teacher", self.ok("accept", stage="1", number=1).get("confirm", ""))

        self.ok("accept", stage="1", number=1, force=True)

        self.assertNotIn((NEW_TEACHER,), {tuple(cell["teachers"]) for _, _, cell in lessons(self.load("answer.json")[MATH_1])})
        self.assertNotIn(NEW_TEACHER, self.course(MATH_1)["scheduled"])

    def test_teacher_paragraph_comes_before_moved_lessons(self):
        """Вариант и переносит урок ``MATH_1``, и меняет ему преподавателя: абзац о смене преподавателя
        (короткий и важный) стоит в вопросе раньше списка переставленных уроков.
        """
        self.openSwapProject()
        saveVariant(self.folder, "1", 1, self.stageOne(**{MATH_1: courseWeek("Математика", NEW_TEACHER, (1, 0), (2, 0))}))

        confirm = self.ok("accept", stage="1", number=1).get("confirm", "")
        teacher, moved = (keyPattern(key).search(confirm) for key in ("web.preview.confirm_teacher", "web.preview.confirm_moved"))

        self.assertTrue(teacher and moved, confirm)
        self.assertLess(teacher.start(), moved.start())


class CopyBlockerAcceptTests(RealProjectCase):
    """Поток 1 ещё не начался. «ЕГЭ основной» Потока 3 присоединяется к Потоку 1, «ЕГЭ продвинутый»
    Потока 3 — к Потоку 2; Поток 2 кончается до начала Потока 3. У ``OLD_TEACHER`` общий урок
    «ЕГЭ продвинутый» Потоков 2 и 3 во вт 1-й урок. Вариант Потока 1 ставит туда ``MATH_1`` того же
    преподавателя — его копия в Потоке 3 мешает копии «Математики» «ЕГЭ продвинутый» Потока 3.
    """
    NAME = "__test_copy_blocker__"

    def test_accept_question_names_source_stage(self):
        """Вопрос принятия: помеха — урок копии в Потоке 3, а составить заново советует Поток 2."""
        settings, answer = jointProject()
        shiftStream(settings, 1, "2026-10-12")
        setSectionEnd(settings, 2, "2027-01-10")
        answer[MATH_LEVEL_2] = courseWeek("Математика", OLD_TEACHER, (1, 0), (4, 1))
        markJoint(settings, 3, JOINT_LINE, 1, answer)
        markJoint(settings, 3, LEVEL_LINE, 2, answer)
        self.openProject(settings, answer)
        saveVariant(self.folder, "1", 1, stageVariant(settings, answer, **{MATH_1: courseWeek("Математика", OLD_TEACHER, (1, 0), (3, 0))}))

        confirm = self.ok("accept", stage="1", number=1).get("confirm", "")
        paragraph = next((part for part in confirm.split("\n\n") if keyPattern("web.preview.confirm_joint_conflicts").search(part)), "")

        self.assertIn(tr("menu.main.tab.classes.slot_busy", detail=jointCourse(3, "Математика", LEVEL_LINE)), paragraph)
        self.assertRegex(paragraph, keyPattern("web.preview.confirm_joint_conflicts", number=3, stages="Поток 2"))


class StaleStartedAcceptTests(RealProjectCase):
    """Поток 1 идёт; у ``MATH_1`` (``OLD_TEACHER``, пн и ср 1-й урок) нагрузку подняли до 3 — курс идёт,
    но ему не хватает урока. Вариант собран до начала курса: у курса другие часы и преподаватель.
    """
    NAME = "__test_stale_started__"

    def test_accept_refuses_even_with_force(self):
        """Принятие такого варианта — отказ web.error.variant_stale_started с курсом и с force,
        расписание не меняется; /variants отдаёт тот же курс в ``staleStarted``.
        """
        settings, answer = jointProject()
        setHours(settings, MATH_1, 3)
        self.openProject(settings, answer)
        saveVariant(self.folder, "1", 1, stageVariant(settings, answer, **{MATH_1: courseWeek("Математика", NEW_TEACHER, (1, 0), (3, 0), (4, 0))}))
        before = self.raw("answer.json")

        self.assertEqual(self.variants("1")["variants"][0]["staleStarted"], [MATH_1])

        for force in (False, True):
            error = self.refused("accept", stage="1", number=1, force=force)
            self.assertEqual(error, tr("web.error.variant_stale_started", courses=f"• {MATH_1}"))

        self.assertEqual(self.raw("answer.json"), before)

    def test_course_started_at_build_is_refused_with_reason(self):
        """Тот же вариант, но курс шёл уже при сборке (started.json: сборка закрепляла его уроки, а решатель
        закрепление не поставил): это не «собран до начала» (``staleStarted`` пуст), но время идущего курса
        не меняется (решение заказчика 06.10.2026) — /variants отдаёт курс в ``movedStarted`` со строками
        «день и час: причина», принятие отказывает и с force (web.error.variant_moved_started), расписание
        не меняется. У ``OLD_TEACHER`` в пн 1-й урок «не может» — эта причина и названа; у урока ср причины
        нет — «программа не смогла оставить урок».
        """
        settings, answer = jointProject()
        setHours(settings, MATH_1, 3)
        markCannot(settings, OLD_TEACHER, "1", (0, 0))
        self.openProject(settings, answer)
        saveVariant(self.folder, "1", 1, stageVariant(settings, answer, **{MATH_1: courseWeek("Математика", NEW_TEACHER, (1, 0), (3, 0), (4, 0))}))
        saveBuildStarted(self.folder, "1", {MATH_1})
        before = self.raw("answer.json")
        lines = [
            f"{slotText(settings, 0, 0)}: {tr('menu.main.tab.classes.slot_unavailable', detail=OLD_TEACHER)}",
            f"{slotText(settings, 2, 0)}: {tr('menu.main.tab.classes.slot_unplaced')}",
        ]

        item = self.variants("1")["variants"][0]
        self.assertEqual(item["staleStarted"], [])
        self.assertEqual(item["movedStarted"], [{"course": MATH_1, "lines": lines}])

        for force in (False, True):
            error = self.refused("accept", stage="1", number=1, force=force)
            # Курс — своей строкой, уроки с причинами — под ним (в полном названии курса уже есть тире)
            self.assertEqual(error, tr("web.error.variant_moved_started", lessons=f"• {MATH_1}" + "".join(f"\n    {line}" for line in lines)))

        self.assertEqual(self.raw("answer.json"), before)

    def test_fresh_variant_of_started_course_is_accepted(self):
        """Курс шёл при сборке, вариант оставил его уроки и преподавателя и добавил недостающий урок —
        ``movedStarted`` пуст, вариант принимается.
        """
        settings, answer = jointProject()
        setHours(settings, MATH_1, 3)
        self.openProject(settings, answer)
        week = courseWeek("Математика", OLD_TEACHER, (0, 0), (2, 0), (4, 0))
        saveVariant(self.folder, "1", 1, stageVariant(settings, answer, **{MATH_1: week}))
        saveBuildStarted(self.folder, "1", {MATH_1})

        self.assertEqual(self.variants("1")["variants"][0]["movedStarted"], [])

        self.ok("accept", stage="1", number=1, force=True)
        self.assertEqual(self.load("answer.json")[MATH_1], week)

    def test_variant_of_older_schedule_is_stale_not_moved(self):
        """Курс шёл при сборке, ему не хватало двух уроков. Вариант 1 добавил один (пт), вариант 2 — два
        других. Приняли вариант 1: курсу всё ещё не хватает урока, принятие берёт его из варианта. Вариант 2
        собран по прежнему расписанию (started.json запомнил уроки курса на начало сборки), урока пт в нём
        нет — это не «программа не смогла оставить урок» (``movedStarted`` пуст), а устаревший вариант:
        ``staleStarted``, отказ web.error.variant_stale_started.
        """
        settings, answer = jointProject()
        setHours(settings, MATH_1, 4)
        self.openProject(settings, answer)
        saveVariant(self.folder, "1", 1, stageVariant(settings, answer, **{MATH_1: courseWeek("Математика", OLD_TEACHER, (0, 0), (2, 0), (4, 0))}))
        saveVariant(self.folder, "1", 2, stageVariant(settings, answer, **{MATH_1: courseWeek("Математика", OLD_TEACHER, (0, 0), (1, 0), (2, 0), (3, 0))}))
        saveBuildStarted(self.folder, "1", {MATH_1}, self.load("answer.json"))

        self.assertEqual([item["movedStarted"] for item in self.variants("1")["variants"]], [[], []])

        self.ok("accept", stage="1", number=1, force=True)
        second = next(item for item in self.variants("1")["variants"] if item["number"] == 2)

        self.assertEqual((second["staleStarted"], second["movedStarted"]), ([MATH_1], []))
        self.assertEqual(self.refused("accept", stage="1", number=2, force=True), tr("web.error.variant_stale_started", courses=f"• {MATH_1}"))


class StageClashAcceptTests(RealProjectCase):
    """Поток 1 идёт, ``MATH_1`` зафиксирован (``OLD_TEACHER``, пн и ср 1-й урок), ``MATH_LEVEL_1`` начнётся
    позже. Вариант собран, пока ``MATH_1`` не шёл: ему — ``NEW_TEACHER``, а ``OLD_TEACHER`` —
    ``MATH_LEVEL_1`` в те же часы. После подмены ``MATH_1`` из расписания у ``OLD_TEACHER`` два урока сразу дважды.
    """
    NAME = "__test_stage_clash__"

    def test_accept_asks_about_clash(self):
        """Метрика teacherClash варианта — 2; принятие без force спрашивает web.preview.confirm_clash."""
        settings, answer = jointProject()
        shiftStream(settings, 1, "2026-09-14")
        next(group for group in settings["classes"]["custom_groups"] if group["name"] == MATH_LEVEL_1)["start_date"] = "2026-10-12"
        self.openProject(settings, answer)
        saveVariant(self.folder, "1", 1, stageVariant(settings, answer, **{
            MATH_1: courseWeek("Математика", NEW_TEACHER, (1, 0), (3, 0)),
            MATH_LEVEL_1: courseWeek("Математика", OLD_TEACHER, (0, 0), (2, 0)),
        }))

        self.assertEqual(self.variants("1")["variants"][0]["metrics"]["teacherClash"], 2)

        confirm = self.ok("accept", stage="1", number=1).get("confirm", "")

        self.assertIn(tr("web.preview.confirm_clash", count=2), confirm)


class LockedVariantTests(RealProjectCase):
    """Вариант Потока 1 собран, пока курсы не шли; теперь Поток 1 идёт и все его уроки на месте."""
    NAME = "__test_locked_variant__"

    def setUp(self):
        super().setUp()
        settings, answer = jointProject()
        settings["max_courses_per_teacher"] = 2
        answer[MATH_OWN_2] = courseWeek("Математика", NEW_TEACHER, (4, 2))
        self.openProject(settings, answer)
        variant = {name: week for name, week in answer.items() if name in stageCourses(settings, "1")}
        variant[MATH_1] = courseWeek("Математика", NEW_TEACHER, (0, 0), (2, 0))
        variant[MATH_LEVEL_1] = courseWeek("Математика", NEW_TEACHER, (1, 0), (3, 0))
        saveVariant(self.folder, "1", 1, variant)

    def test_variant_shows_answer_teacher(self):
        """«Кто ведёт» у зафиксированного ``MATH_1`` — преподаватель из расписания: принятие его не сменит."""
        item = self.variants("1")["variants"][0]

        self.assertEqual(item["teachers"][MATH_1], [OLD_TEACHER])
        self.assertEqual(item["answer"][MATH_1], self.load("answer.json")[MATH_1])

    def test_accept_does_not_count_locked_courses(self):
        """Лимит 2: по файлу варианта у ``NEW_TEACHER`` было бы 3 курса, но два из них зафиксированы
        с другими преподавателями — вопроса о лимите нет.
        """
        body = self.ok("accept", stage="1", number=1)

        self.assertNoText("web.preview.confirm_over_limit", body.get("confirm", ""))
